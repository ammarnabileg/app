import os
from docx import Document

input_path = r'c:\Users\Administrator\Desktop\zahed\عقد حكومي بدل سكن.docx'
output_path = r'c:\Users\Administrator\Desktop\newhr\contract_template.docx'

doc = Document(input_path)

replacements = {
    'يونغ ميلينغ': '{{ employee_name }}',
    'ماليزيا': '{{ nationality }}',
    '286101408992': '{{ civil_id }}',
    'مراقب مالي': '{{ profession }}',
    '3500': '{{ salary }}',
    '02/08/2026': '{{ start_date }}',
    'YONG MEE LING': '{{ employee_name_en }}',
    'Malaysia': '{{ nationality_en }}',
    'Financial Controller': '{{ profession_en }}'
}

def replace_text_in_runs(paragraphs):
    for p in paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                # To preserve formatting, we have to find the runs that contain the text.
                # However, python-docx splits text across runs unpredictably.
                # A simple approach for templates is to replace the text of the first run 
                # and clear the subsequent runs if the text spans multiple.
                # But to keep it simple and robust, we can just clear all runs and put the replaced text in the first run, 
                # OR we just replace the text in paragraph.text (which loses run-level formatting if it varied).
                pass
        
        # Simple paragraph level replacement (might lose some inline styling like bolding parts of a word)
        original_text = p.text
        new_text = original_text
        for key, val in replacements.items():
            new_text = new_text.replace(key, val)
        
        if new_text != original_text:
            # We clear runs and add back
            p.clear()
            p.add_run(new_text)

replace_text_in_runs(doc.paragraphs)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            replace_text_in_runs(cell.paragraphs)

doc.save(output_path)
print(f"Template saved to {output_path}")
