"""اعتمادُ الساعات إلى Excel — وصفحةُ الاعتماد نفسُها.

## ما طُلب

صفٌّ لكلّ موظّف فيه أيّامُه كلُّها، والموظّفون تحت بعض بالشكل نفسه:
رقم الملف، الاسم، القسم، المسمّى، أيّام العمل، المرضيّة/الطارئة، عمودٌ لكلّ
يوم («20/08خميس»)، ثم المطلوبة والفعليّة والناقصة والإضافيّ والإجمالي
والتأخير والحضور والغياب، ثم الأساسيّ وقيمةُ الإضافيّ والخصومات والصافي.

## ما وُجد في الطريق

صفحةُ الاعتماد كانت تسقط (500) منذ 2.16.3.0: نصٌّ فيه `%(n)s` يُمرَّر إلى
`_()` في القالب بلا معاملات، وgettext في القوالب يطبّق `%` دائمًا فيرمي
KeyError. والاختبارُ كان يفتح الطباعةَ لا الصفحة. والآن فحصٌ لكلّ القوالب.

يُشغَّل:  python -m pytest tests/test_hours_export.py -v
"""
import glob
import io
import os
import re
import sqlite3
import sys
from datetime import date, timedelta

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

MONTH, YEAR = 8, 2026


@pytest.fixture
def web(tmp_path, monkeypatch):
    import importlib
    monkeypatch.setenv('HR_DATA_DIR', str(tmp_path))
    import utils.db as db
    importlib.reload(db)
    db.init_db()
    conn = sqlite3.connect(os.path.join(str(tmp_path), 'hr_system.db'))
    conn.row_factory = sqlite3.Row
    conn.execute("INSERT INTO shift_types (name, start_time, end_time, hours_per_day, flex_mode, is_active)"
                 " VALUES ('S', '08:00', '16:00', 8, 'none', 1)")
    for eid, dept in ((1, 'الإدارة'), (2, 'المخزن')):
        conn.execute("INSERT INTO employees (id, employee_number, name, arabic_name, department, position,"
                     " hire_date, salary, is_active, shift_type) VALUES (?, ?, ?, ?, ?, 'فنّي',"
                     " '2020-01-01', 260, 1, 'S')", (eid, str(eid), f'E{eid}', f'موظف {eid}', dept))
    conn.commit()
    import utils.auth as auth
    importlib.reload(auth)
    auth.get_current_license_info = lambda: {'ok': True}
    import app as A
    importlib.reload(A)
    A.app.before_request_funcs[None] = [f for f in A.app.before_request_funcs.get(None, [])
                                        if f.__name__ not in ('check_license_globally',
                                                              'enforce_plan_features')]
    A.app.config['TESTING'] = True
    admin = conn.execute("SELECT id FROM users WHERE role='admin'").fetchone()[0]
    c = A.app.test_client()
    with c.session_transaction() as s:
        s.update({'user_id': admin, 'role': 'admin', 'username': 'admin'})
    import utils.payroll_engine as pe
    return c, conn, A, pe


def _punch(conn, eid, day, *times):
    for t in times:
        conn.execute("INSERT INTO attendance_records (employee_id, device_id, check_time, check_type)"
                     " VALUES (?, 'T1', ?, 'I')", (eid, f'{day} {t}'))
    conn.commit()


def _sheet(c, **q):
    from openpyxl import load_workbook
    qs = '&'.join(f'{k}={v}' for k, v in {'month': MONTH, 'year': YEAR, **q}.items())
    r = c.get('/payroll/hours_approval/export?' + qs)
    assert r.status_code == 200, r.status_code
    assert r.mimetype.endswith('spreadsheetml.sheet')
    return load_workbook(io.BytesIO(r.data)).active


def _header(ws):
    return [ws.cell(4, i).value for i in range(1, ws.max_column + 1)]


def _rows(ws):
    """صفوفُ الموظّفين: من تحت العناوين حتى صفّ المجاميع."""
    out = {}
    for r in range(5, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None:
            break
        out[str(v)] = [ws.cell(r, i).value for i in range(1, ws.max_column + 1)]
    return out


def _hours(v):
    """openpyxl يقرأ الوقت timedelta؛ ونُعيده ساعاتٍ عشريّة للمقارنة."""
    return v.total_seconds() / 3600 if isinstance(v, timedelta) else v


def test_one_row_per_employee_with_every_day_of_the_period(web):
    c, conn, A, pe = web
    start, end = pe.resolve_period(conn, MONTH, YEAR)
    n = (end - start).days + 1
    ws = _sheet(c)
    head = _header(ws)
    assert head[:6] == ['رقم الملف', 'الاسم', 'القسم', 'المسمى الوظيفي', 'أيام العمل', 'المرضية/الطارئة']
    assert head[6] == start.strftime('%d/%m') + 'السبت'          # 1 أغسطس 2026 سبت
    assert len(head) == 6 + n + 12
    assert head[6 + n:] == ['الساعات المطلوبة', 'الساعات الفعلية', 'الساعات الناقصة', 'إضافي',
                            'الإجمالي', 'التأخير', 'الحضور', 'الغياب', 'الراتب الأساسي',
                            'قيمة الإضافي', 'إجمالي الخصومات', 'صافي الراتب']
    assert set(_rows(ws)) == {'1', '2'}
    assert ws.sheet_view.rightToLeft and ws.freeze_panes == 'G5'


def test_day_cells_show_hours_or_a_code(web):
    c, conn, A, pe = web
    start, _ = pe.resolve_period(conn, MONTH, YEAR)
    sun = start + timedelta(days=(7 - start.isoweekday() % 7) % 7)   # أوّل أحد
    fri = sun + timedelta(days=5)
    _punch(conn, 1, sun.isoformat(), '08:00', '16:51')             # 8:51 ساعة
    _punch(conn, 1, (sun + timedelta(days=1)).isoformat(), '08:00')  # بصمةٌ ناقصة
    _punch(conn, 1, fri.isoformat(), '09:00', '13:00')             # جمعةٌ اشتُغلت
    sick = conn.execute("SELECT id FROM leave_types WHERE law_kind = 'sick'").fetchone()[0]
    conn.execute("INSERT INTO leave_requests (employee_id, leave_type_id, start_date, end_date, days_count,"
                 " status, is_paid_leave) VALUES (1, ?, ?, ?, 1, 'approved', 1)",
                 (sick, (sun + timedelta(days=2)).isoformat(), (sun + timedelta(days=2)).isoformat()))
    conn.commit()
    ws = _sheet(c)
    row = _rows(ws)['1']
    col = lambda d: 6 + (d - start).days                            # noqa: E731
    assert _hours(row[col(sun)]) == pytest.approx(8 + 51 / 60)
    assert ws.cell(5, col(sun) + 1).number_format == '[h]:mm'
    assert row[col(sun + timedelta(days=1))] == 'ن'
    assert row[col(sun + timedelta(days=2))] == 'م'
    assert row[col(sun + timedelta(days=3))] == 'غ'
    assert _hours(row[col(fri)]) == pytest.approx(4.0)
    assert row[col(fri + timedelta(days=1))] == 'ر'                 # السبتُ راحة
    assert row[5] == 1, 'المرضية/الطارئة'


def test_the_totals_match_the_engine(web):
    c, conn, A, pe = web
    start, _ = pe.resolve_period(conn, MONTH, YEAR)
    sun = start + timedelta(days=(7 - start.isoweekday() % 7) % 7)
    _punch(conn, 1, sun.isoformat(), '08:20', '18:00')              # تأخير 20 د + إضافيّ ساعتان
    ws = _sheet(c)
    head = _header(ws)
    row = dict(zip(head, _rows(ws)['1']))
    m = pe.compute_month_metrics(conn, MONTH, YEAR)[1]
    pay = next(r for r in pe.compute_monthly_payroll(conn, MONTH, YEAR)['rows'] if r['employee_id'] == 1)
    assert row['أيام العمل'] == m['required_days']
    assert _hours(row['الساعات المطلوبة']) == pytest.approx(m['required_hours'])
    assert _hours(row['الساعات الفعلية']) == pytest.approx(m['actual_hours'])
    # الإضافيُّ لا يسدّ نقصًا: الناقصةُ من العاديّة (الفعليّة − الإضافيّ)
    assert _hours(row['الساعات الناقصة']) == pytest.approx(m['required_hours'] - (m['actual_hours'] - 2))
    assert _hours(row['إضافي']) == pytest.approx(2.0)
    assert _hours(row['الإجمالي']) == pytest.approx(m['actual_hours']), 'لا يُعدّ الإضافيّ مرّتين'
    assert _hours(row['التأخير']) == pytest.approx(20 / 60)
    assert (row['الحضور'], row['الغياب']) == (m['actual_days'], m['absent_days'])
    assert row['صافي الراتب'] == pytest.approx(pay['net'])
    assert row['إجمالي الخصومات'] == pytest.approx(pay['total_deductions'])
    ot_item = next(a for a in pay['allowances'] if a['code'] == 'overtime')
    assert row['قيمة الإضافي'] == pytest.approx(ot_item['amount'])
    assert ot_item['amount'] == pytest.approx(2 * 1.25 * 1.25, abs=0.01)


def test_the_totals_row_sums_the_employees(web):
    c, conn, A, pe = web
    ws = _sheet(c)
    head = _header(ws)
    rows = _rows(ws)
    total_row = 5 + len(rows)
    assert ws.cell(total_row, 2).value == 'الإجمالي'
    i = head.index('الغياب') + 1
    assert ws.cell(total_row, i).value == sum(r[i - 1] for r in rows.values())
    j = head.index('صافي الراتب') + 1
    assert ws.cell(total_row, j).value == pytest.approx(sum(r[j - 1] for r in rows.values()))


def test_the_department_filter_applies(web):
    c, conn, A, pe = web
    assert set(_rows(_sheet(c, dept='المخزن'))) == {'2'}


def test_no_salary_figures_without_the_payroll_permission(web):
    c, conn, A, pe = web
    conn.execute("INSERT INTO roles (name) VALUES ('معتمد ساعات')")
    rid = conn.execute("SELECT id FROM roles WHERE name = 'معتمد ساعات'").fetchone()[0]
    for code in ('page.hours_approval', 'salary.approve_hours'):
        conn.execute('INSERT INTO role_permissions (role_id, permission_id) '
                     'SELECT ?, id FROM permissions WHERE code = ?', (rid, code))
    conn.execute("INSERT INTO users (username, password, full_name, role, is_active)"
                 " VALUES ('hr', 'x', 'hr', 'user', 1)")
    uid = conn.execute("SELECT id FROM users WHERE username = 'hr'").fetchone()[0]
    conn.execute('INSERT INTO user_roles (user_id, role_id) VALUES (?, ?)', (uid, rid))
    conn.commit()
    h = A.app.test_client()
    with h.session_transaction() as s:
        s.update({'user_id': uid, 'role': 'user', 'username': 'hr'})
    head = _header(_sheet(h))
    assert head[-1] == 'الغياب'
    assert 'صافي الراتب' not in head and 'الراتب الأساسي' not in head


def test_a_bad_period_is_refused(web):
    c, conn, A, pe = web
    assert c.get('/payroll/hours_approval/export?month=13&year=2026').status_code == 400
    assert c.get('/payroll/hours_approval/export?month=x&year=2026').status_code == 400


# ------------------------------------------------ الصفحة

def test_the_hours_approval_page_opens(web):
    c, conn, A, pe = web
    r = c.get(f'/payroll/hours_approval?month={MONTH}&year={YEAR}')
    body = r.get_data(as_text=True)
    assert r.status_code == 200, r.status_code
    assert 'id="haExportBtn"' in body
    assert '%(n)s' in body and '%(policy)s' in body, 'العلاماتُ تصل JavaScript كما هي'


def _po(lang):
    s = open(os.path.join(ROOT, 'translations', lang, 'LC_MESSAGES', 'messages.po'), encoding='utf-8').read()
    return dict(re.findall(r'^msgid "(.*?)"\nmsgstr "(.*?)"$', s, re.M))


def test_no_template_passes_a_format_string_to_gettext_without_values():
    """`_('مفتاح')` في القالب يُطبَّق عليه `% {}` دائمًا: `%(n)s` يرمي KeyError
    ويُسقط الصفحة، و`%s` يصير «{}». والعلامةُ الحرفيّة تُكتب `%%`."""
    cats = {lang: _po(lang) for lang in ('ar', 'en')}
    bad = []
    for f in glob.glob(os.path.join(ROOT, 'templates', '**', '*.html'), recursive=True):
        text = open(f, encoding='utf-8').read()
        for m in re.finditer(r"_\(\s*'([^']+)'\s*\)", text):
            for lang, cat in cats.items():
                if '%' in (cat.get(m.group(1)) or '').replace('%%', ''):
                    bad.append((os.path.relpath(f, ROOT), m.group(1), lang))
    assert bad == [], bad


def test_an_incomplete_split_shift_day_is_a_missing_punch_not_hours(web):
    c, conn, A, pe = web
    conn.execute("INSERT INTO shift_types (name, start_time, end_time, hours_per_day, flex_mode, is_active,"
                 " is_split) VALUES ('SP', '08:00', '21:00', 8, 'none', 1, 1)")
    sid = conn.execute("SELECT id FROM shift_types WHERE name = 'SP'").fetchone()[0]
    conn.execute("INSERT INTO shift_periods (shift_type_id, seq, start_time, end_time) VALUES (?, 1, '08:00', '12:00')", (sid,))
    conn.execute("INSERT INTO shift_periods (shift_type_id, seq, start_time, end_time) VALUES (?, 2, '17:00', '21:00')", (sid,))
    conn.execute("UPDATE employees SET shift_type = 'SP' WHERE id = 2")
    conn.commit()
    start, _ = pe.resolve_period(conn, MONTH, YEAR)
    sun = start + timedelta(days=(7 - start.isoweekday() % 7) % 7)
    _punch(conn, 2, sun.isoformat(), '08:00', '12:00')          # الفترةُ الثانية بلا بصمة
    day = next(d for d in pe.compute_employee_days(conn, 2, MONTH, YEAR)['days'] if d['date'] == sun.isoformat())
    assert day['status'] == 'partial' and day['hours'] > 0
    assert _rows(_sheet(c))['2'][6 + (sun - start).days] == 'ن'
