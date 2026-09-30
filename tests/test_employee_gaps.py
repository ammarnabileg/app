"""استكمالُ البيانات الناقصة وطريقةُ الدفع (utils/employee_gaps).

ما يظهر بالأحمر في كشف الرواتب يُستكمل من ملفّ Excel واحد: يُنزَّل بالخانات
الناقصة حمراء، ويُملأ، ويُرفع — وكلُّ تغييرٍ في سجلّ الموظّف. والراتبُ يُملأ هنا
لمن راتبه صفر فقط؛ تغييرُ راتبٍ قائم زيادةٌ لها تاريخ سريان من شاشة الموظّف.

يُشغَّل:  python -m pytest tests/test_employee_gaps.py -v
"""
import io
import os
import sqlite3
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)


@pytest.fixture
def web(tmp_path, monkeypatch):
    import importlib
    monkeypatch.setenv('HR_DATA_DIR', str(tmp_path))
    import utils.db as db
    importlib.reload(db)
    db.init_db()
    conn = sqlite3.connect(os.path.join(str(tmp_path), 'hr_system.db'))
    conn.row_factory = sqlite3.Row
    rows = (
        # رقم، عربي، تعيين، راتب، حساب، مركز، موقع، عقد، طريقة
        (1, 'موظف كامل', '2020-01-01', 300, 'KW1', 'CC', 'HQ', 'full_time', 'bank'),
        (2, '', '2021-05-01', 250, '', '', '', '', None),            # ناقص: عربي، مركز، موقع، عقد، دفع
        (3, 'بلا راتب', '2022-01-01', 0, '', 'CC', 'HQ', 'full_time', 'cash'),
        (4, 'تحويل بلا حساب', '2022-01-01', 400, '', 'CC', 'HQ', 'full_time', 'bank'),
    )
    for eid, ar, hire, sal, acct, cc, loc, ct, pm in rows:
        conn.execute("INSERT INTO employees (id, employee_number, name, arabic_name, department, position,"
                     " hire_date, salary, is_active, bank_account_number, cost_center, branch_location,"
                     " contract_type) VALUES (?, ?, ?, ?, 'D', 'P', ?, ?, 1, ?, ?, ?, ?)",
                     (eid, str(eid), f'E{eid}', ar, hire, sal, acct, cc, loc, ct))
    conn.commit()
    from utils.labor_law import migrate
    migrate(conn)
    for eid, *_r, pm in rows:
        conn.execute('UPDATE employees SET payment_method = ? WHERE id = ?', (pm, eid))
    conn.commit()
    import utils.auth as auth
    importlib.reload(auth)
    auth.get_current_license_info = lambda: {'ok': True}
    import app as A
    importlib.reload(A)
    A.app.before_request_funcs[None] = [f for f in A.app.before_request_funcs.get(None, [])
                                        if f.__name__ not in ('check_license_globally', 'enforce_plan_features')]
    A.app.config['TESTING'] = True
    admin = conn.execute("SELECT id FROM users WHERE role='admin'").fetchone()[0]
    c = A.app.test_client()
    with c.session_transaction() as s:
        s.update({'user_id': admin, 'role': 'admin', 'username': 'admin'})
    return c, conn, admin


def _emp(conn, eid):
    return conn.execute('SELECT * FROM employees WHERE id = ?', (eid,)).fetchone()


# ------------------------------------------------ ما الناقص

def test_payment_label_and_gaps():
    from utils.employee_gaps import payment_label, gaps
    assert payment_label({'payment_method': 'bank', 'bank_account_number': 'KW1'}) == ('Bank Transfer', True)
    assert payment_label({'payment_method': 'bank', 'bank_account_number': ''}) == ('Bank Transfer', False)
    assert payment_label({'payment_method': 'check'}) == ('Check', True)
    assert payment_label({'payment_method': 'cash'}) == ('Cash', True)
    assert payment_label({'payment_method': None, 'bank_iban': 'KW81'}) == ('Bank Transfer', True)
    assert payment_label({'payment_method': None}) == ('Cash', False)
    base = {'arabic_name': 'a', 'hire_date': '2020-01-01', 'department': 'D', 'position': 'P',
            'cost_center': 'C', 'branch_location': 'L', 'contract_type': 'full_time', 'salary': 100,
            'payment_method': 'cash'}
    assert gaps(base) == {}
    assert gaps(dict(base, salary=0)) == {'salary': 'missing'}
    assert gaps(dict(base, payment_method='bank')) == {'payment_method': 'missing', 'bank_account_number': 'missing'}
    assert gaps(base, shared={'2020-01-01'}) == {'hire_date': 'review'}


def test_the_page_summarises_and_the_file_marks_missing_cells_red(web):
    from openpyxl import load_workbook
    from utils.employee_gaps import RED, FIELDS
    c, conn, _admin = web
    body = c.get('/employees/data_gaps').get_data(as_text=True)
    assert 'id="gapsSummary"' in body and 'id="gapsDownload"' in body
    assert 'id="dataGapsBtn"' in c.get('/employees').get_data(as_text=True)
    r = c.get('/employees/data_gaps/export')
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.data)).worksheets[0]
    head = [x.value for x in ws[1]]
    rows = {str(ws.cell(i, 1).value): i for i in range(2, ws.max_row + 1)}
    assert set(rows) == {'2', '3', '4'}, 'الموظّفُ كاملُ البيانات لا يظهر'
    col = {h: head.index(h) + 1 for h in head}
    red = lambda r, h: (ws.cell(r, col[h]).fill.fgColor.rgb or '') == RED
    r2, r3, r4 = rows['2'], rows['3'], rows['4']
    assert red(r2, 'الاسم العربي') and red(r2, 'مركز التكلفة') and red(r2, 'طريقة الدفع')
    assert not red(r2, 'الراتب الأساسي') and ws.cell(r2, col['الراتب الأساسي']).value == 250
    assert red(r3, 'الراتب الأساسي') and not red(r3, 'طريقة الدفع')
    assert red(r4, 'رقم الحساب') and ws.cell(r4, col['طريقة الدفع']).value == 'Bank Transfer'
    assert len(head) == 2 + len(FIELDS)


def _upload(c, fill):
    """ينزّل الملف، يملؤه بـ fill {رقم: {عنوان: قيمة}}، ويرفعه."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get('/employees/data_gaps/export').data))
    ws = wb.worksheets[0]
    head = [x.value for x in ws[1]]
    for i in range(2, ws.max_row + 1):
        for h, v in fill.get(str(ws.cell(i, 1).value), {}).items():
            ws.cell(i, head.index(h) + 1).value = v
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return c.post('/employees/data_gaps', data={'file': (buf, 'gaps.xlsx')},
                  content_type='multipart/form-data', follow_redirects=True)


def test_upload_applies_filled_cells_and_logs_them(web):
    c, conn, admin = web
    r = _upload(c, {
        '2': {'الاسم العربي': 'اسم عربي', 'مركز التكلفة': 'CC9', 'الموقع': 'Fahaheel',
              'نوع العقد': 'open_ended', 'طريقة الدفع': 'Check'},
        '4': {'رقم الحساب': 'KW99NBOK1'},
        '3': {'مركز التكلفة': None, 'الموقع': ''},          # خانةٌ مُسحت لا تمحو القيمة المسجّلة
    })
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and 'id="gapsResult"' in body
    e2, e4 = _emp(conn, 2), _emp(conn, 4)
    assert (e2['arabic_name'], e2['cost_center'], e2['branch_location'], e2['contract_type'],
            e2['payment_method']) == ('اسم عربي', 'CC9', 'Fahaheel', 'open_ended', 'check')
    assert e4['bank_account_number'] == 'KW99NBOK1'
    assert (_emp(conn, 3)['cost_center'], _emp(conn, 3)['branch_location']) == ('CC', 'HQ')
    log = {(x['employee_id'], x['field']) for x in conn.execute(
        "SELECT employee_id, field FROM employee_audit_log WHERE source = 'data_gaps'")}
    assert {(2, 'arabic_name'), (2, 'cost_center'), (2, 'payment_method'), (4, 'bank_account_number')} <= log
    from utils.employee_gaps import employees_with_gaps
    assert {e['employee_number'] for e, _g in employees_with_gaps(conn)} == {'3'}


def test_salary_is_filled_only_where_it_is_missing(web):
    c, conn, _admin = web
    body = _upload(c, {'3': {'الراتب الأساسي': 275}, '2': {'الراتب الأساسي': 999}}).get_data(as_text=True)
    assert _emp(conn, 3)['salary'] == 275
    hist = [tuple(x) for x in conn.execute('SELECT effective_from, salary FROM employee_salary_history '
                                           'WHERE employee_id = 3')]
    assert hist == [('0001-01-01', 275.0)], 'راتبٌ كان صفرًا يصير الراتبَ من البداية'
    assert _emp(conn, 2)['salary'] == 250, 'راتبٌ قائم لا يتغيّر من الملف'
    assert 'شاشة الموظف' in body


def test_bad_values_are_reported_and_skipped(web):
    c, conn, _admin = web
    body = _upload(c, {'2': {'تاريخ التعيين': '2099-01-01', 'طريقة الدفع': 'Bitcoin',
                             'نوع العقد': 'forever', 'الاسم العربي': 'صحيح'}}).get_data(as_text=True)
    e2 = _emp(conn, 2)
    assert e2['hire_date'] == '2021-05-01' and e2['payment_method'] is None and e2['contract_type'] == ''
    assert e2['arabic_name'] == 'صحيح', 'الخانة الصحيحة تُطبَّق ولو أخطأت جارتُها'
    for msg in ('تاريخ تعيين في المستقبل', 'طريقة دفع غير معروفة', 'نوع عقد غير معروف'):
        assert msg in body


def test_a_foreign_file_is_refused(web):
    from openpyxl import Workbook
    c, conn, _admin = web
    wb = Workbook()
    wb.active.append(['x', 'y'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    def flashes(data):
        r = c.post('/employees/data_gaps', data=data, content_type='multipart/form-data')
        assert r.status_code == 302
        with c.session_transaction() as s:
            return [m for _cat, m in s.pop('_flashes', [])]
    assert any('رقم الملف' in m for m in flashes({'file': (buf, 'x.xlsx')}))
    assert flashes({'file': (io.BytesIO(b'x'), 'x.csv')}) == ['اختر ملف Excel بصيغة ‎.xlsx']
    assert _emp(conn, 2)['arabic_name'] == ''


def test_the_edit_form_saves_the_payment_method(web):
    c, conn, _admin = web
    body = c.get('/employees/edit/2').get_data(as_text=True)
    assert 'name="payment_method"' in body
    from test_payroll_engine_v2 import _edit_form
    c.post('/employees/edit/2', data=_edit_form(conn, payment_method='check', eid=2))
    assert _emp(conn, 2)['payment_method'] == 'check'
    c.post('/employees/edit/2', data=_edit_form(conn, payment_method='', eid=2))
    assert _emp(conn, 2)['payment_method'] is None


def test_payment_method_drives_the_sheet_and_the_bank_file(web):
    from utils.payroll_sheet import missing_cells
    from utils.payslip import bank_workbook
    from openpyxl import load_workbook
    c, conn, _admin = web
    conn.execute("UPDATE employees SET payment_method = 'check' WHERE id = 2")
    conn.commit()
    from utils.payroll_sheet import build_sheet
    rows = {r['emp_no']: r for r in build_sheet(conn, 8, 2026)['rows']}
    assert rows['1']['pay_type'] == 'Bank Transfer' and 'AQ' not in missing_cells(rows['1'])
    assert rows['2']['pay_type'] == 'Check' and 'AQ' not in missing_cells(rows['2'])
    assert rows['3']['pay_type'] == 'Cash' and 'AQ' not in missing_cells(rows['3'])
    assert rows['4']['pay_type'] == 'Bank Transfer' and 'AQ' in missing_cells(rows['4'])
    # يُصرف لمن له أيّامٌ مستحقّة
    for eid in (1, 2, 3, 4):
        for d in range(3, 8):
            for t in ('08:00', '16:00'):
                conn.execute("INSERT INTO attendance_records (employee_id, device_id, check_time, check_type)"
                             " VALUES (?, 'T', ?, 'I')", (eid, f'2026-08-{d:02d} {t}'))
    conn.execute("UPDATE employees SET salary = 100 WHERE id = 3")
    conn.execute("UPDATE employee_salary_history SET salary = 100 WHERE employee_id = 3")
    conn.commit()
    buf, _src = bank_workbook(conn, 8, 2026, 'X')
    wb = load_workbook(buf)
    ids = {name: [str(wb[name].cell(r, 1).value) for r in range(4, wb[name].max_row + 1)
                  if wb[name].cell(r, 6).value and wb[name].cell(r, 2).value != 'Total']
           for name in ('Bank', 'Cash', 'Check')}
    assert ids['Bank'] == ['1'] and ids['Check'] == ['2'] and sorted(ids['Cash']) == ['3', '4']
