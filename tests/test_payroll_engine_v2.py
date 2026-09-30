"""محرّكُ الرواتب 2.22 — تاريخُ الشفت والراتب، والليليّ، والغياب، والقسيمة.

## ما سُدّ

- **تاريخُ الشفتات:** كلُّ يومٍ بالشفت الساري فيه. موظّفٌ صباحيّ 15 يومًا ثم
  مسائيّ 5 ثم صباحيّ يُحسب كلُّ يومٍ بشفته — ولا يتغيّر الماضي بتغيير اليوم.
- **الشفتُ الليليّ:** يملك بصماتِ صباح الغد (كان يُقسم عند منتصف الليل).
- **المقسّم:** الشهرُ يجمع الفترات لا المدى بينها (كما يعرض اليوم).
- **تاريخُ الراتب:** زيادةٌ يوم 16 تُحسب من 16. والبدلُ الذي يبدأ في وسط الشهر
  بقدر أيّامه.
- **الأجرُ اليوميّ:** الراتبُ أجرُ يوم، ويعادل شهريًّا × القاسم.
- **الغياب:** بالأجر كلّه (إعداد)، ولا أجرَ بلا يومٍ مستحقّ، ولا صافيَ سالبًا.
- **القسيمة، وملفُّ البنك، والمخصّصات.**

يُشغَّل:  python -m pytest tests/test_payroll_engine_v2.py -v
"""
import io
import json
import os
import sqlite3
import sys
from datetime import date, timedelta

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

MONTH, YEAR = 8, 2026           # أغسطس 2026: الدورة 1–31، يبدأ بالسبت


@pytest.fixture
def env(tmp_path, monkeypatch):
    import importlib
    monkeypatch.setenv('HR_DATA_DIR', str(tmp_path))
    import utils.db as db
    importlib.reload(db)
    db.init_db()
    conn = sqlite3.connect(os.path.join(str(tmp_path), 'hr_system.db'))
    conn.row_factory = sqlite3.Row
    for name, a, b in (('صباح', '08:00', '16:00'), ('مساء', '14:00', '22:00'),
                       ('ليل', '22:00', '06:00')):
        conn.execute("INSERT INTO shift_types (name, start_time, end_time, hours_per_day, flex_mode,"
                     " is_active) VALUES (?, ?, ?, 8, 'none', 1)", (name, a, b))
    _emp(conn, 1, 'صباح')
    conn.commit()
    from utils.labor_law import migrate
    migrate(conn)          # كالإقلاع: الأساسُ لكلّ موظّفٍ قائم
    conn.commit()
    import utils.payroll_engine as pe
    importlib.reload(pe)
    return conn, pe


def _emp(conn, eid, shift, salary=260, hire='2020-01-01', **kw):
    cols = {'id': eid, 'employee_number': str(eid), 'name': f'E{eid}', 'arabic_name': f'موظف {eid}',
            'department': 'D', 'position': 'P', 'hire_date': hire, 'salary': salary,
            'is_active': 1, 'shift_type': shift}
    cols.update(kw)
    conn.execute(f"INSERT INTO employees ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                 list(cols.values()))


def _set(conn, **kw):
    for k, v in kw.items():
        conn.execute('INSERT INTO salary_settings_v2 (setting_name, setting_value) VALUES (?, ?) '
                     'ON CONFLICT(setting_name) DO UPDATE SET setting_value = excluded.setting_value',
                     (k, str(v)))
    conn.commit()


def _punch(conn, eid, *stamps):
    for s in stamps:
        conn.execute("INSERT INTO attendance_records (employee_id, device_id, check_time, check_type)"
                     " VALUES (?, 'T', ?, 'I')", (eid, s))
    conn.commit()


def _days(pe, conn, eid=1):
    return {d['date']: d for d in pe.compute_employee_days(conn, eid, MONTH, YEAR)['days']}


def _row(pe, conn, eid=1):
    return next(r for r in pe.compute_monthly_payroll(conn, MONTH, YEAR)['rows'] if r['employee_id'] == eid)


def _item(items, code):
    return next((i for i in items if i['code'] == code), None)


def _workdays():
    d, out = date(2026, 8, 1), []
    while d.month == 8:
        if d.isoweekday() not in (5, 6):
            out.append(d.isoformat())
        d += timedelta(days=1)
    return out


# ================================================== تاريخُ الشفتات

def _rotation(conn):
    """صباحيّ حتى 15، مسائيّ 16–20، صباحيّ من 21 — مثالُ السؤال."""
    from utils.workday import record_shift_change
    record_shift_change(conn, 1, 'مساء', '2026-08-16', previous='صباح')
    record_shift_change(conn, 1, 'صباح', '2026-08-21')
    conn.commit()


def test_each_day_uses_the_shift_in_effect_that_day(env):
    conn, pe = env
    _rotation(conn)
    _punch(conn, 1, '2026-08-13 08:00', '2026-08-13 16:00',     # صباحيّ
           '2026-08-17 14:00', '2026-08-17 22:00',              # مسائيّ
           '2026-08-18 14:30', '2026-08-18 22:00',              # مسائيّ متأخّر 30 د
           '2026-08-23 08:00', '2026-08-23 16:00')              # صباحيّ من جديد
    d = _days(pe, conn)
    assert (d['2026-08-13']['shift_name'], d['2026-08-17']['shift_name'],
            d['2026-08-23']['shift_name']) == ('صباح', 'مساء', 'صباح')
    assert d['2026-08-17']['late_mins'] == 0 and d['2026-08-17']['ot_mins'] == 0
    assert d['2026-08-18']['late_mins'] == 30
    m = pe.compute_month_metrics(conn, MONTH, YEAR)[1]
    assert m['late_mins'] == 30, 'بلا التاريخ: مسائيٌّ يُقرأ صباحيًّا فيُحسب تأخير 6 ساعات'
    assert m['ot_weekday_mins'] == 0, 'وخروجُ 22:00 لا يُقرأ إضافيًّا'


def test_changing_the_shift_today_keeps_past_days(env):
    conn, pe = env
    _punch(conn, 1, '2026-08-04 08:00', '2026-08-04 16:00')
    before = pe.compute_month_metrics(conn, MONTH, YEAR)[1]
    from utils.workday import record_shift_change
    record_shift_change(conn, 1, 'مساء', '2026-09-01', previous='صباح')
    conn.execute("UPDATE employees SET shift_type = 'مساء' WHERE id = 1")
    conn.commit()
    after = pe.compute_month_metrics(conn, MONTH, YEAR)[1]
    assert (after['late_mins'], after['early_mins'], after['required_hours']) == \
        (before['late_mins'], before['early_mins'], before['required_hours'])


def test_a_new_employee_gets_a_baseline_on_the_first_change(env):
    conn, pe = env
    _emp(conn, 2, 'صباح')
    conn.commit()
    assert not conn.execute('SELECT 1 FROM employee_shift_history WHERE employee_id = 2').fetchone()
    from utils.workday import record_shift_change
    record_shift_change(conn, 2, 'مساء', '2026-08-16', previous='صباح')
    conn.execute("UPDATE employees SET shift_type = 'مساء' WHERE id = 2")
    conn.commit()
    d = _days(pe, conn, 2)
    assert d['2026-08-10']['shift_name'] == 'صباح' and d['2026-08-20']['shift_name'] == 'مساء'


def test_the_migration_seeds_the_current_shift_from_the_start(env):
    conn, pe = env
    row = conn.execute('SELECT shift_type, effective_from FROM employee_shift_history WHERE employee_id = 1').fetchone()
    assert (row[0], row[1]) == ('صباح', '0001-01-01')


# ================================================== الشفتُ الليليّ والمقسّم

def test_a_night_shift_owns_the_next_mornings_punches(env):
    conn, pe = env
    _emp(conn, 2, 'ليل')
    conn.commit()
    _punch(conn, 2, '2026-08-02 22:10', '2026-08-03 06:30', '2026-08-03 21:55', '2026-08-04 05:40')
    d = _days(pe, conn, 2)
    n1, n2 = d['2026-08-02'], d['2026-08-03']
    assert (n1['status'], n1['check_in'], n1['check_out'], n1['hours']) == ('present', '22:10', '06:30', 8.33)
    assert (n2['status'], n2['check_in'], n2['check_out']) == ('present', '21:55', '05:40')
    assert n1['late_mins'] == 0, 'عشرُ دقائق داخل السماح'
    assert n1['ot_mins'] == 30 and n2['early_mins'] == 20
    m = pe.compute_month_metrics(conn, MONTH, YEAR)[2]
    assert m['partial_days'] == 0 and m['actual_hours'] == pytest.approx(8.33 + 7.75, abs=0.01)


def test_a_late_night_arrival_after_midnight_is_late(env):
    conn, pe = env
    _emp(conn, 2, 'ليل')
    conn.commit()
    _punch(conn, 2, '2026-08-05 00:30', '2026-08-05 06:00')      # ليلةُ 4 أغسطس، وصل 00:30
    d = _days(pe, conn, 2)
    assert d['2026-08-04']['status'] == 'present' and d['2026-08-04']['late_mins'] == 150


def test_night_then_morning_does_not_steal_the_morning_checkin(env):
    conn, pe = env
    from utils.workday import record_shift_change
    record_shift_change(conn, 1, 'ليل', '2026-08-10', previous='صباح')
    record_shift_change(conn, 1, 'صباح', '2026-08-11')
    conn.commit()
    _punch(conn, 1, '2026-08-10 22:00', '2026-08-11 06:00', '2026-08-11 08:05', '2026-08-11 16:00')
    d = _days(pe, conn)
    assert (d['2026-08-10']['check_in'], d['2026-08-10']['check_out']) == ('22:00', '06:00')
    assert (d['2026-08-11']['check_in'], d['2026-08-11']['check_out']) == ('08:05', '16:00')


def test_split_shift_month_hours_sum_the_periods(env):
    conn, pe = env
    conn.execute("INSERT INTO shift_types (name, start_time, end_time, hours_per_day, flex_mode, is_active,"
                 " is_split) VALUES ('مقسم', '08:00', '21:00', 8, 'none', 1, 1)")
    sid = conn.execute("SELECT id FROM shift_types WHERE name = 'مقسم'").fetchone()[0]
    conn.execute("INSERT INTO shift_periods (shift_type_id, seq, start_time, end_time) VALUES (?, 1, '08:00', '12:00')", (sid,))
    conn.execute("INSERT INTO shift_periods (shift_type_id, seq, start_time, end_time) VALUES (?, 2, '17:00', '21:00')", (sid,))
    _emp(conn, 2, 'مقسم')
    conn.commit()
    _punch(conn, 2, '2026-08-02 08:00', '2026-08-02 12:00', '2026-08-02 17:00', '2026-08-02 21:00',
           '2026-08-03 08:00', '2026-08-03 12:00')                 # الثاني: فترةٌ واحدة
    m = pe.compute_month_metrics(conn, MONTH, YEAR)[2]
    assert m['actual_hours'] == pytest.approx(12.0), '8 + 4 — لا 13 ولا 17'
    assert m['partial_days'] == 1 and m['actual_days'] == 1
    assert _days(pe, conn, 2)['2026-08-03']['status'] == 'partial'


# ================================================== تاريخُ الراتب والبدلات

def test_a_raise_mid_month_counts_from_its_day(env):
    conn, pe = env
    from utils.pay_history import record_salary_change
    record_salary_change(conn, 1, 400, '2026-08-16', previous=260)
    conn.commit()
    r = _row(pe, conn)
    assert r['basic'] == pytest.approx(round((15 * 260 + 16 * 400) / 31, 3), abs=0.001)
    assert [p['salary'] for p in r['salary_parts']] == [260, 400]
    assert [p['days'] for p in r['salary_parts']] == [15, 16]


def test_an_unchanged_salary_is_exact(env):
    conn, pe = env
    assert _row(pe, conn)['basic'] == 260


def test_an_allowance_starting_mid_month_counts_its_days(env):
    conn, pe = env
    tid = conn.execute("SELECT id FROM payroll_item_types WHERE code = 'housing'").fetchone()[0]
    conn.execute("INSERT INTO employee_salary_items (employee_id, item_type_id, calc_method, amount,"
                 " is_active, start_date) VALUES (1, ?, 'fixed', 62, 1, '2026-08-11')", (tid,))
    conn.commit()
    h = _item(_row(pe, conn)['allowances'], 'housing')
    assert h['amount'] == pytest.approx(62 - 62 * 10 / 31, abs=0.001)
    assert h['full_amount'] == 62 and h['item_missing_days'] == 10


def test_daily_paid_salary_is_a_daily_wage(env):
    conn, pe = env
    _emp(conn, 2, 'صباح', salary=10, pay_type='daily')
    conn.commit()
    for d in _workdays()[:5]:
        _punch(conn, 2, f'{d} 08:00', f'{d} 16:00')
    r = _row(pe, conn, 2)
    assert r['basic'] == 260 and r['daily_rate'] == pytest.approx(10)
    ab = _item(r['deductions'], 'absence_deduction')
    assert ab['amount'] == pytest.approx(ab['days'] * 10)


# ================================================== الغياب والصافي

def _housing(conn, eid=1, amt=104):
    tid = conn.execute("SELECT id FROM payroll_item_types WHERE code = 'housing'").fetchone()[0]
    conn.execute("INSERT INTO employee_salary_items (employee_id, item_type_id, calc_method, amount,"
                 " is_active) VALUES (?, ?, 'fixed', ?, 1)", (eid, tid, amt))
    conn.commit()


@pytest.mark.parametrize('basis', ['26', '30'])
def test_a_fully_absent_month_pays_nothing_and_never_goes_negative(env, basis):
    conn, pe = env
    _housing(conn)
    _set(conn, daily_rate_basis=basis)
    r = _row(pe, conn)
    assert r['net'] == 0
    assert 'nothing_earned' in {w['code'] for w in r['law_warnings']}


def test_absence_on_the_whole_wage_or_the_basic(env):
    conn, pe = env
    _housing(conn)
    days = _workdays()
    for d in days[1:]:
        _punch(conn, 1, f'{d} 08:00', f'{d} 16:00')                 # غاب يومًا واحدًا
    _set(conn, absence_deduction_base='wage')
    assert _item(_row(pe, conn)['deductions'], 'absence_deduction')['amount'] == pytest.approx(14)
    _set(conn, absence_deduction_base='basic')
    assert _item(_row(pe, conn)['deductions'], 'absence_deduction')['amount'] == pytest.approx(10)


def test_a_month_of_missing_punches_is_not_a_month_of_absence(env):
    conn, pe = env
    for d in _workdays()[:10]:
        _punch(conn, 1, f'{d} 08:00')                                # حضر، ونسي الخروج — والباقي غياب
    r = _row(pe, conn)
    assert _item(r['deductions'], 'absence_deduction')['days'] == len(_workdays()) - 10
    assert 'nothing_earned' not in {w['code'] for w in r['law_warnings']}
    assert r['net'] > 0


def test_the_net_never_goes_below_zero(env):
    conn, pe = env
    for d in _workdays():
        _punch(conn, 1, f'{d} 08:00', f'{d} 16:00')
    tid = conn.execute("SELECT id FROM payroll_item_types WHERE code = 'damages'").fetchone()[0]
    conn.execute("INSERT INTO payroll_transactions (employee_id, item_type_id, month, year, amount, status)"
                 " VALUES (1, ?, ?, ?, 400, 'active')", (tid, MONTH, YEAR))
    conn.commit()
    r = _row(pe, conn)
    assert r['net'] == 0
    w = next(w for w in r['law_warnings'] if w['code'] == 'net_floor')
    assert w['amount'] == pytest.approx(140)
    assert _item(r['deductions'], 'damages')['deferred'] == pytest.approx(140)


def test_new_installs_use_the_whole_wage_and_three_decimals(env):
    conn, pe = env
    from utils.settings_utils import get_salary_settings_v2
    s = get_salary_settings_v2(conn)
    assert (s['absence_deduction_base'], s['rounding_display_decimals']) == ('wage', '3')


def test_existing_installs_keep_the_basic_until_they_choose(tmp_path, monkeypatch):
    import importlib
    monkeypatch.setenv('HR_DATA_DIR', str(tmp_path))
    import utils.db as db
    importlib.reload(db)
    db.init_db()
    conn = sqlite3.connect(os.path.join(str(tmp_path), 'hr_system.db'))
    conn.row_factory = sqlite3.Row
    conn.execute("DELETE FROM salary_settings_v2 WHERE setting_name = 'absence_deduction_base'")
    _emp(conn, 9, 'x')
    conn.commit()
    from utils.labor_law import migrate
    migrate(conn)
    assert conn.execute("SELECT setting_value FROM salary_settings_v2 WHERE setting_name = "
                        "'absence_deduction_base'").fetchone()[0] == 'basic'


# ================================================== عبر الشاشات

@pytest.fixture
def web(env):
    import importlib
    conn, pe = env
    import utils.auth as auth
    importlib.reload(auth)
    auth.get_current_license_info = lambda: {'ok': True}
    import app as A
    importlib.reload(A)
    A.app.before_request_funcs[None] = [f for f in A.app.before_request_funcs.get(None, [])
                                        if f.__name__ not in ('check_license_globally',
                                                              'enforce_plan_features')]
    A.app.config['TESTING'] = True
    admin = conn.execute("SELECT id FROM users WHERE role='admin'").fetchone()[0]
    c = A.app.test_client()
    with c.session_transaction() as s:
        s.update({'user_id': admin, 'role': 'admin', 'username': 'admin'})
    return c, conn, A, pe


def _edit_form(conn, eid=1, **kw):
    emp = dict(conn.execute('SELECT * FROM employees WHERE id = ?', (eid,)).fetchone())
    form = {k: ('' if v is None else str(v)) for k, v in emp.items()}
    form.update(kw)
    return form


def test_editing_the_shift_records_its_effective_date(web):
    c, conn, A, pe = web
    c.post('/employees/edit/1', data=_edit_form(conn, shift_type='مساء',
                                                shift_effective_from='2026-08-16'))
    hist = conn.execute('SELECT shift_type, effective_from FROM employee_shift_history '
                        'WHERE employee_id = 1 ORDER BY effective_from').fetchall()
    assert [(h[0], h[1]) for h in hist] == [('صباح', '0001-01-01'), ('مساء', '2026-08-16')]
    d = _days(pe, conn)
    assert d['2026-08-10']['shift_name'] == 'صباح' and d['2026-08-20']['shift_name'] == 'مساء'


def test_editing_the_salary_records_its_effective_date(web):
    c, conn, A, pe = web
    c.post('/employees/edit/1', data=_edit_form(conn, salary='390', salary_effective_from='2026-08-16'))
    r = _row(pe, conn)
    assert r['basic'] == pytest.approx(round((15 * 260 + 16 * 390) / 31, 3), abs=0.001)
    assert conn.execute('SELECT salary FROM employees WHERE id = 1').fetchone()[0] == 390


def test_a_future_change_does_not_apply_before_its_day(web):
    c, conn, A, pe = web
    future = (date.today() + timedelta(days=30)).isoformat()
    c.post('/employees/edit/1', data=_edit_form(conn, shift_type='مساء', shift_effective_from=future))
    assert conn.execute('SELECT shift_type FROM employees WHERE id = 1').fetchone()[0] == 'صباح'


def test_history_inside_a_locked_period_is_refused(web):
    c, conn, A, pe = web
    conn.execute("INSERT INTO payroll_runs (month, year, status, employees_count) VALUES (?, ?, 'approved', 1)",
                 (MONTH, YEAR))
    conn.commit()
    c.post('/employees/1/shift_history', data={'shift_type': 'مساء', 'effective_from': '2026-08-16'})
    c.post('/employees/edit/1', data=_edit_form(conn, salary='999', salary_effective_from='2026-08-10'))
    assert conn.execute('SELECT COUNT(*) FROM employee_shift_history WHERE employee_id = 1').fetchone()[0] == 1
    assert conn.execute('SELECT salary FROM employees WHERE id = 1').fetchone()[0] == 260
    hid = conn.execute('SELECT id FROM employee_shift_history WHERE employee_id = 1').fetchone()[0]
    conn.execute("UPDATE employee_shift_history SET effective_from = '2026-08-05' WHERE id = ?", (hid,))
    conn.commit()
    c.post(f'/employees/1/shift_history/{hid}/delete')
    assert conn.execute('SELECT COUNT(*) FROM employee_shift_history WHERE id = ?', (hid,)).fetchone()[0] == 1


def test_history_routes_add_and_delete(web):
    c, conn, A, pe = web
    c.post('/employees/1/shift_history', data={'shift_type': 'مساء', 'effective_from': '2026-08-16'})
    c.post('/employees/1/salary_history', data={'salary': '300', 'effective_from': '2026-08-16'})
    assert conn.execute('SELECT COUNT(*) FROM employee_shift_history WHERE employee_id = 1').fetchone()[0] == 2
    assert conn.execute('SELECT COUNT(*) FROM employee_salary_history WHERE employee_id = 1').fetchone()[0] == 2
    c.post('/employees/1/shift_history', data={'shift_type': 'لا يوجد', 'effective_from': '2026-08-20'})
    assert conn.execute('SELECT COUNT(*) FROM employee_shift_history WHERE employee_id = 1').fetchone()[0] == 2
    hid = conn.execute("SELECT id FROM employee_salary_history WHERE effective_from = '2026-08-16'").fetchone()[0]
    c.post(f'/employees/1/salary_history/{hid}/delete')
    assert conn.execute('SELECT COUNT(*) FROM employee_salary_history WHERE employee_id = 1').fetchone()[0] == 1
    body = c.get('/employees/edit/1').get_data(as_text=True)
    assert 'id="shiftHistTable"' in body and 'id="histShiftForm"' in body and 'مساء' in body


def test_one_click_fix_applies_the_compliant_value(web):
    c, conn, A, pe = web
    _set(conn, overtime_round_to_minutes=15, absence_deduction_base='basic', rounding_display_decimals=2)
    body = c.get('/settings/labor-law').get_data(as_text=True)
    assert '/settings/labor-law/fix/absence_base' in body
    for code in ('ot_rounding', 'absence_base', 'decimals'):
        c.post(f'/settings/labor-law/fix/{code}')
    from utils.settings_utils import get_salary_settings_v2
    s = get_salary_settings_v2(conn)
    assert (s['overtime_round_to_minutes'], s['absence_deduction_base'],
            s['rounding_display_decimals']) == ('0', 'wage', '3')


# ------------------------------------------------ القسيمة، البنك، المخصّصات

def test_the_admin_payslip_shows_the_engine_figures(web):
    c, conn, A, pe = web
    for d in _workdays()[:10]:
        _punch(conn, 1, f'{d} 08:00', f'{d} 16:00')
    net = _row(pe, conn)['net']
    body = c.get(f'/payroll/payslips?month={MONTH}&year={YEAR}&employee_id=1').get_data(as_text=True)
    assert 'قسيمة الراتب' in body and '%.3f' % net in body and 'data-emp="1"' in body
    assert 'مسودة' in body, 'قبل الحفظ: مسودة'


def _portal_client(conn, A, eid):
    conn.execute("INSERT INTO users (username, password, full_name, role, is_active, employee_id)"
                 " VALUES (?, 'x', 'e', 'user', 1, ?)", (f'u{eid}', eid))
    conn.commit()
    uid = conn.execute('SELECT id FROM users WHERE username = ?', (f'u{eid}',)).fetchone()[0]
    p = A.app.test_client()
    with p.session_transaction() as s:
        s.update({'user_id': uid, 'role': 'user', 'employee_id': eid})
    return p


def test_the_portal_payslip_is_the_saved_one_and_only_ones_own(web):
    c, conn, A, pe = web
    _emp(conn, 2, 'صباح', salary=500)
    conn.commit()
    p = _portal_client(conn, A, 1)
    body = p.get(f'/portal/payslip?month={MONTH}&year={YEAR}').get_data(as_text=True)
    assert 'data-emp=' not in body, 'قبل الحفظ لا قسيمة'
    conn.execute("INSERT INTO payroll_runs (month, year, status, employees_count, period_start, period_end)"
                 " VALUES (?, ?, 'saved', 2, '2026-08-01', '2026-08-31')", (MONTH, YEAR))
    rid = conn.execute('SELECT id FROM payroll_runs').fetchone()[0]
    for eid, net in ((1, 111.5), (2, 999.0)):
        conn.execute("INSERT INTO payroll_run_lines (run_id, employee_id, basic, net, details_json)"
                     " VALUES (?, ?, 260, ?, ?)", (rid, eid, net, json.dumps({'allowances': [], 'deductions': []})))
    conn.commit()
    body = p.get(f'/portal/payslip?month={MONTH}&year={YEAR}&employee_id=2').get_data(as_text=True)
    assert 'data-emp="1"' in body and '111.500' in body
    assert '999.000' not in body, 'لا يرى الموظّفُ قسيمةَ غيره'


def test_the_bank_file_splits_bank_and_cash(web):
    from openpyxl import load_workbook
    c, conn, A, pe = web
    _emp(conn, 2, 'صباح', salary=500, bank_name='NBK', bank_account_number='KW81NBOK0000')
    conn.commit()
    for eid in (1, 2):
        for d in _workdays():
            _punch(conn, eid, f'{d} 08:00', f'{d} 16:00')
    r = c.get(f'/payroll/monthly/bank_export?month={MONTH}&year={YEAR}')
    wb = load_workbook(io.BytesIO(r.data))
    bank, cash = wb['Bank'], wb['Cash']
    assert 'NOT SAVED' in bank.cell(1, 1).value
    assert bank.cell(4, 5).value == 'KW81NBOK0000' and bank.cell(4, 6).value == pytest.approx(500)
    assert cash.cell(4, 1).value == '1' and cash.cell(4, 6).value == pytest.approx(260)


def test_provisions_match_the_end_of_service_calculation(web):
    from openpyxl import load_workbook
    from routes.eos_routes import compute_kuwait_eos
    c, conn, A, pe = web
    r = c.get('/payroll/provisions/export?as_of=2026-08-31')
    ws = load_workbook(io.BytesIO(r.data)).active
    calc = compute_kuwait_eos(conn, 1, '2026-08-31', 'termination')
    assert ws.cell(5, 1).value == '1'
    assert ws.cell(5, 8).value == pytest.approx(calc['gross_gratuity'])
    assert ws.cell(5, 11).value == pytest.approx(calc['gross_gratuity'] + calc['leave_amount'], abs=0.001)
    assert c.get('/payroll/provisions/export?as_of=bad').status_code == 400
