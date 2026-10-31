# -*- coding: utf-8 -*-
"""Excel التقرير الشامل المخصّص: منسّق من السيرفر، والساعات وقت حقيقي بيتجمع، وسطر إجمالي."""
import io
from datetime import timedelta

import pytest

from tests.test_hours_export import MONTH, YEAR, _punch, web  # noqa: F401


def _h(v):
    return v.total_seconds() / 3600 if isinstance(v, timedelta) else v * 24


def test_export_is_styled_and_matches_the_page(web):
    c, conn, A, pe = web
    from openpyxl import load_workbook
    import importlib
    import smart_reports
    importlib.reload(smart_reports)            # DB_PATH بيتاخد وقت الـ import — من قاعدة الاختبار اللي قبله
    start, end = pe.resolve_period(conn, MONTH, YEAR)
    sun = start + timedelta(days=(7 - start.isoweekday() % 7) % 7)
    _punch(conn, 1, sun.isoformat(), '08:00', '16:30')
    q = f'start_date={start.isoformat()}&end_date={end.isoformat()}'
    page = c.get('/api/fingerprint/custom_employees_report?' + q).get_json()
    r = c.get('/fingerprint/custom_employees_report/export?' + q)
    assert r.status_code == 200 and r.mimetype.endswith('spreadsheetml.sheet')
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb.worksheets[0]
    assert ws.sheet_view.rightToLeft and ws.freeze_panes == 'C7' and ws.print_title_rows == '$5:$6'
    assert ws.page_setup.orientation == 'landscape'
    days = len(page['rows'][0]['daily_cells'])
    assert ws.cell(5, 7).value == f'{start:%d/%m}' and ws.cell(6, 7).value          # التاريخ واليوم
    head5 = [ws.cell(5, i).value for i in range(1, ws.max_column + 1)]
    net_col = len(head5)
    rows = {ws.cell(r, 1).value: r for r in range(7, ws.max_row + 1)}
    r1 = rows[1]
    col = 6 + (sun - start).days + 1
    assert ws.cell(r1, col).number_format == '[h]:mm' and abs(_h(ws.cell(r1, col).value) - 8.5) < 1e-6
    assert ws.cell(r1, net_col).value == round(page['rows'][0]['financials']['net_salary'], 3)
    total = ws.cell(ws.max_row, net_col).value
    assert str(total).startswith('=SUM('), 'سطر إجمالي بمعادلة'
    assert any(ws.cell(r1, 6 + i + 1).fill.fgColor.rgb.endswith('FBD5D5') for i in range(days)), 'الغياب ملوّن'
    det = wb.worksheets[1]
    assert det.auto_filter.ref and det.cell(3, 6).value
    assert det.max_row == 3 + days * len(page['rows'])


def test_inactive_employee_with_punches_shows_and_zero_days_means_zero_salary(web):
    c, conn, A, pe = web
    import importlib
    import smart_reports
    importlib.reload(smart_reports)
    start, end = pe.resolve_period(conn, MONTH, YEAR)
    _punch(conn, 2, (start + timedelta(days=2)).isoformat(), '08:00', '16:00')
    conn.execute('UPDATE employees SET is_active = 0 WHERE id = 2')        # موقوف من غير تاريخ نهاية خدمة
    conn.commit()
    q = f'start_date={start.isoformat()}&end_date={end.isoformat()}'
    rows = {r['id']: r for r in c.get('/api/fingerprint/custom_employees_report?' + q).get_json()['rows']}
    assert 2 in rows and rows[2]['stats']['present_days'] == 1, 'الموقوف بأيامه يظهر'
    assert rows[1]['stats']['present_days'] == 0
    assert rows[1]['financials']['net_salary'] == 0, 'مداومش ولا يوم = مالوش راتب'
    dept = rows[2]['department']
    only = c.get('/api/fingerprint/custom_employees_report?' + q + f'&dept={dept}').get_json()['rows']
    assert {r['department'] for r in only} == {dept}, 'فلتر القسم على الكل'


def test_report_numbers_are_the_payroll_numbers(web):
    """الحالة اللي اتبعتت: يومين بصمة ناقصة وقعدة بعد الشفت — التقرير كان بيقول إضافي صفر ومفيش خصم."""
    c, conn, A, pe = web
    import importlib
    import smart_reports
    importlib.reload(smart_reports)
    start, end = pe.resolve_period(conn, MONTH, YEAR)
    d, work, partial = start, 0, 0
    while d <= end:
        if d.isoweekday() not in (5, 6):
            if partial < 2:
                _punch(conn, 1, d.isoformat(), '08:00')                     # بصمة ناقصة
                partial += 1
            else:
                _punch(conn, 1, d.isoformat(), '08:00', '17:00')            # ساعة بعد الشفت
            work += 1
        d += timedelta(days=1)
    q = f'start_date={start.isoformat()}&end_date={end.isoformat()}'
    row = next(r for r in c.get('/api/fingerprint/custom_employees_report?' + q).get_json()['rows'] if r['id'] == 1)
    pay = next(r for r in pe.compute_monthly_payroll(conn, MONTH, YEAR)['rows'] if r['employee_id'] == 1)
    met = pe.compute_month_metrics(conn, MONTH, YEAR)[1]
    assert row['financials']['source'] == 'payroll_engine' and not row['financials']['partial_period']
    assert row['financials']['net_salary'] == pytest.approx(pay['net'], abs=0.001), 'نفس صافي كشف الرواتب'
    assert row['financials']['total_deduction'] == pytest.approx(pay['total_deductions'], abs=0.001)
    assert row['stats']['overtime_hours'] * 60 == pytest.approx(met['ot_weekday_mins']) and met['ot_weekday_mins'] > 0
    assert row['stats']['required_hours'] == pytest.approx(met['required_hours'])
    # فترة جزئية: نص الدورة ← الأساسي بالنسبة
    mid = start + timedelta(days=14)
    half = next(r for r in c.get(f'/api/fingerprint/custom_employees_report?start_date={start}&end_date={mid}').get_json()['rows']
                if r['id'] == 1)
    assert half['financials']['partial_period'] and half['salary'] < pay['basic']
