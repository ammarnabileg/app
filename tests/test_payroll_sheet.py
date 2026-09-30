"""كشفُ الرواتب بشكل شيت الشركة (Monthly Payroll) — Excel وطباعة.

الشيتُ الذي تعمل به الشركة: الراتبُ بالعقد، ثم المستحقُّ عن أيّام العمل، ثم
الإضافات (الإضافيّ بساعاته ومبالغه — يوميّ وجمعة وعطلة — والإجازة والعمولة
وأخرى)، ثم الإجماليّ، ثم الخصومات (غياب، جزاءات، سلف، أخرى)، ثم الصافي.

كلُّ مبلغٍ بندٌ من محرّك الرواتب. والحارسُ الأوّل هنا: **صافي الشيت = صافي
المحرّك** لكلّ موظّف، وإلّا فالشيتُ يقول غيرَ ما يُصرف.

يُشغَّل:  python -m pytest tests/test_payroll_sheet.py -v
"""
import io
import os
import sqlite3
import sys
from datetime import date, timedelta

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

MONTH, YEAR = 8, 2026


@pytest.fixture
def env(tmp_path, monkeypatch):
    import importlib
    monkeypatch.setenv('HR_DATA_DIR', str(tmp_path))
    import utils.db as db
    importlib.reload(db)
    db.init_db()
    conn = sqlite3.connect(os.path.join(str(tmp_path), 'hr_system.db'))
    conn.row_factory = sqlite3.Row
    conn.execute("INSERT INTO shift_types (name, start_time, end_time, hours_per_day, flex_mode, is_active)"
                 " VALUES ('S', '08:00', '16:00', 8, 'none', 1)")
    emps = ((1, '2020-01-01', 260, 'KW123'), (2, '2026-08-10', 520, ''))   # الثاني عُيّن في منتصف الشهر
    for eid, hire, sal, bank in emps:
        conn.execute("INSERT INTO employees (id, employee_number, name, arabic_name, department, position,"
                     " hire_date, salary, is_active, shift_type, bank_account_number, contract_type,"
                     " cost_center, branch_location) VALUES (?, ?, ?, ?, 'HQ.', 'Clerk', ?, ?, 1, 'S', ?,"
                     " 'full_time', 'CC1', 'Kuwait')",
                     (eid, str(eid), f'EMP {eid}', f'موظف {eid}', hire, sal, bank))

    def item(code):
        return conn.execute('SELECT id FROM payroll_item_types WHERE code = ?', (code,)).fetchone()[0]
    for code, amt in (('transport', 30), ('housing', 50), ('phone', 5), ('hazard', 15)):
        conn.execute("INSERT INTO employee_salary_items (employee_id, item_type_id, calc_method, amount,"
                     " is_active) VALUES (1, ?, 'fixed', ?, 1)", (item(code), amt))
    for code, amt in (('commission', 12.5), ('incentive', 7), ('admin_penalty', 4), ('damages', 3)):
        conn.execute("INSERT INTO payroll_transactions (employee_id, item_type_id, month, year, amount,"
                     " status) VALUES (1, ?, ?, ?, ?, 'active')", (item(code), MONTH, YEAR, amt))
    conn.execute("INSERT INTO employee_loans (employee_id, principal, monthly_installment, start_month,"
                 " start_year, status) VALUES (1, 500, 20, 1, 2026, 'active')")
    conn.commit()
    import utils.payroll_engine as pe
    importlib.reload(pe)
    start, _ = pe.resolve_period(conn, MONTH, YEAR)
    sun = start + timedelta(days=(7 - start.isoweekday() % 7) % 7)
    punches = [(sun, '08:00', '18:00'),                       # إضافيّ ساعتان
               (sun + timedelta(days=1), '08:20', '16:00'),   # تأخير 20 د
               (sun + timedelta(days=5), '09:00', '12:30')]   # جمعة: 3:30
    for d, a, b in punches:
        for t in (a, b):
            conn.execute("INSERT INTO attendance_records (employee_id, device_id, check_time, check_type)"
                         " VALUES (1, 'T', ?, 'I')", (f'{d.isoformat()} {t}',))
    conn.commit()
    return conn, pe


def _sheet(conn):
    from utils.payroll_sheet import build_sheet
    return {r['emp_no']: r for r in build_sheet(conn, MONTH, YEAR)['rows']}, build_sheet(conn, MONTH, YEAR)


def _engine(pe, conn, eid):
    return next(r for r in pe.compute_monthly_payroll(conn, MONTH, YEAR)['rows'] if r['employee_id'] == eid)


def test_the_sheet_net_is_the_engine_net(env):
    conn, pe = env
    rows, _ = _sheet(conn)
    for eid in (1, 2):
        assert rows[str(eid)]['net'] == pytest.approx(_engine(pe, conn, eid)['net'], abs=1e-6)


def test_contract_salary_columns(env):
    conn, pe = env
    r = _sheet(conn)[0]['1']
    assert (r['c_basic'], r['c_trans'], r['c_house'], r['c_phone'], r['c_other']) == (260, 30, 50, 5, 15)
    assert r['c_allow'] == 100 and r['c_total'] == 360
    assert (r['e_trans'], r['e_house'], r['e_phone'], r['e_other']) == (30, 50, 5, 15)


def test_working_days_prorate_the_basic_like_the_engine(env):
    conn, pe = env
    r = _sheet(conn)[0]['2']
    pre = next(d for d in _engine(pe, conn, 2)['deductions'] if d['code'] == 'pre_hire_days')
    assert r['work_days'] == 26 - pre['days']
    assert r['e_basic'] == pytest.approx(520 - pre['amount'])
    assert r['other_ded'] == pytest.approx(
        sum(d['amount'] for d in _engine(pe, conn, 2)['deductions']
            if d['code'] not in ('pre_hire_days', 'absence_deduction', 'loan_installment')))
    assert _sheet(conn)[0]['1']['work_days'] == 26


def test_overtime_is_split_by_kind_and_adds_up(env):
    conn, pe = env
    r = _sheet(conn)[0]['1']
    assert r['ot_h_daily'] == pytest.approx(2.0) and r['ot_h_fri'] == pytest.approx(3.5)
    assert r['ot_a_daily'] == pytest.approx(2 * 1.25 * 1.25, abs=0.01)
    assert r['ot_a_fri'] == pytest.approx(3.5 * 1.5 * 1.25, abs=0.01)
    assert r['ot_a_daily'] + r['ot_a_fri'] + r['ot_a_hol'] == pytest.approx(r['ot_total'], abs=1e-9)
    ot = next(a for a in _engine(pe, conn, 1)['allowances'] if a['code'] == 'overtime')
    assert r['ot_total'] == pytest.approx(ot['amount'])


def test_additions_gross_and_deductions(env):
    conn, pe = env
    r = _sheet(conn)[0]['1']
    assert r['commission'] == 12.5 and r['others'] == 7
    assert r['gross'] == pytest.approx(r['e_total'] + r['ot_total'] + r['leave'] + r['commission'] + r['others'])
    eng = _engine(pe, conn, 1)
    late = next(d for d in eng['deductions'] if d['code'] == 'lateness_deduction')
    assert r['fine_amt'] == pytest.approx(late['amount'] + 4), 'التأخيرُ والجزاءُ الإداريّ'
    assert r['fine_days'] == pytest.approx(r['fine_amt'] / 10, abs=0.01)
    absent = next(d for d in eng['deductions'] if d['code'] == 'absence_deduction')
    assert (r['abs_days'], r['abs_amt']) == (absent['days'], absent['amount'])
    assert r['loan'] == 20
    assert r['other_ded'] == 3, 'خصمُ العهدة في «أخرى»'
    assert r['total_ded'] == pytest.approx(r['abs_amt'] + r['fine_amt'] + r['loan'] + r['other_ded'])


def test_payment_type_term_and_leave_balance(env):
    conn, pe = env
    rows = _sheet(conn)[0]
    assert rows['1']['pay_type'] == 'Bank Transfer' and rows['2']['pay_type'] == 'Cash'
    assert rows['1']['term'] == 'Full time'
    from utils.leave_balance import compute_leave_balance
    end = pe.resolve_period(conn, MONTH, YEAR)[1]
    assert rows['1']['leave_balance'] == compute_leave_balance(conn, 1, as_of=end)['balance']


def test_totals_sum_the_rows(env):
    conn, pe = env
    rows, sheet = _sheet(conn)
    for k in ('net', 'gross', 'total_ded', 'work_days', 'ot_h_daily'):
        assert sheet['totals'][k] == pytest.approx(sum(r[k] for r in rows.values()), abs=1e-6)


# ------------------------------------------------ عبر الشاشة

@pytest.fixture
def web(env):
    import importlib
    conn, pe = env
    import utils.auth as auth
    importlib.reload(auth)
    auth.get_current_license_info = lambda: {'ok': True}
    import app as A
    importlib.reload(A)
    A.app.before_request_funcs[None] = [f for f in A.app.before_request_funcs.get(None, [])
                                        if f.__name__ not in ('check_license_globally',
                                                              'enforce_plan_features')]
    A.app.config['TESTING'] = True
    admin = conn.execute("SELECT id FROM users WHERE role='admin'").fetchone()[0]
    c = A.app.test_client()
    with c.session_transaction() as s:
        s.update({'user_id': admin, 'role': 'admin', 'username': 'admin'})
    return c, conn, A


def test_the_excel_file_has_the_sheet_layout(web):
    from openpyxl import load_workbook
    c, conn, A = web
    r = c.get(f'/payroll/monthly/sheet/export?month={MONTH}&year={YEAR}')
    assert r.status_code == 200 and r.mimetype.endswith('spreadsheetml.sheet')
    ws = load_workbook(io.BytesIO(r.data)).active
    head = [ws.cell(4, i).value for i in range(1, 9)]
    assert head == ['Emp#.', 'Name_E', 'Name_A', 'DOJ.', 'Cost Center', 'Dept.', 'Location', 'Job']
    assert ws.cell(2, 1).value == 'Monthly Payroll'
    groups = {ws.cell(rng.min_row, rng.min_col).value for rng in ws.merged_cells.ranges}
    assert {'additions', 'Hours', 'Amt.', 'Deduction', 'Absent', 'Fine', 'Gross salary',
            'Net Salary'} <= groups
    from utils.payroll_sheet import KEY_INDEX
    row1 = next(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value == '1')
    assert ws.cell(row1, KEY_INDEX['ot_h_fri']).number_format == '[h]:mm'
    assert ws.cell(row1, KEY_INDEX['ot_h_fri']).value == timedelta(hours=3, minutes=30)
    assert ws.cell(row1, KEY_INDEX['net']).value == pytest.approx(
        next(x for x in _sheet(conn)[0].values() if x['emp_no'] == '1')['net'])
    total_row = row1 + 2
    assert ws.cell(total_row, KEY_INDEX['name_e']).value == 'Total'
    assert ws.freeze_panes == 'I5' and ws.print_title_rows == '$2:$4'
    assert ws.page_setup.orientation == 'landscape'
    assert 'Prepared By' in [ws.cell(ws.max_row, i).value for i in range(1, ws.max_column + 1)]


def test_the_print_view(web):
    c, conn, A = web
    r = c.get(f'/payroll/monthly/sheet/print?month={MONTH}&year={YEAR}')
    body = r.get_data(as_text=True)
    assert r.status_code == 200
    for text in ('Monthly Payroll', 'Emp#.', 'additions', 'Deduction', 'Net Salary', 'Prepared By',
                 'EMP 1', 'موظف 2', '3:30', 'Bank Transfer'):
        assert text in body, text
    net = next(x for x in _sheet(conn)[0].values() if x['emp_no'] == '1')['net']
    assert '{:,.3f}'.format(net) in body
    assert "addEventListener('beforeprint'" in body


def test_bad_periods_and_permissions(web):
    c, conn, A = web
    for u in ('/payroll/monthly/sheet/export', '/payroll/monthly/sheet/print'):
        assert c.get(u + '?month=13&year=2026').status_code == 400
        assert c.get(u + '?month=x&year=2026').status_code == 400
    conn.execute("INSERT INTO users (username, password, full_name, role, is_active) VALUES ('u', 'x', 'u', 'user', 1)")
    conn.commit()
    uid = conn.execute("SELECT id FROM users WHERE username = 'u'").fetchone()[0]
    h = A.app.test_client()
    with h.session_transaction() as s:
        s.update({'user_id': uid, 'role': 'user', 'username': 'u'})
    for u in ('/payroll/monthly/sheet/export', '/payroll/monthly/sheet/print'):
        assert h.get(f'{u}?month={MONTH}&year={YEAR}').status_code in (302, 403)


def test_the_payroll_page_buttons_use_the_sheet(web):
    c, conn, A = web
    body = c.get(f'/payroll/monthly?month={MONTH}&year={YEAR}').get_data(as_text=True)
    assert '/payroll/monthly/sheet/print' in body and '/payroll/monthly/sheet/export' in body
    assert 'pmExportDetailed' in body and '/payroll/monthly/export' in body


# ------------------------------------------------ تناسبُ البدلات مع أيّام الخدمة

def _add_fixed(conn, eid, code, amt):
    tid = conn.execute('SELECT id FROM payroll_item_types WHERE code = ?', (code,)).fetchone()[0]
    conn.execute("INSERT INTO employee_salary_items (employee_id, item_type_id, calc_method, amount,"
                 " is_active) VALUES (?, ?, 'fixed', ?, 1)", (eid, tid, amt))
    conn.commit()


def test_a_mid_month_hire_gets_allowances_for_his_days(env):
    conn, pe = env
    _add_fixed(conn, 2, 'transport', 26)
    _add_fixed(conn, 2, 'housing', 52)
    tid = conn.execute("SELECT id FROM payroll_item_types WHERE code = 'commission'").fetchone()[0]
    conn.execute("INSERT INTO payroll_transactions (employee_id, item_type_id, month, year, amount, status)"
                 " VALUES (2, ?, ?, ?, 40, 'active')", (tid, MONTH, YEAR))
    conn.commit()
    eng = _engine(pe, conn, 2)
    assert next(a for a in eng['allowances'] if a['code'] == 'commission')['amount'] == 40, \
        'العمولةُ ليست بدلًا ثابتًا: لا تُناسَب'
    pre = next(d for d in eng['deductions'] if d['code'] == 'pre_hire_days')['days']
    tr = next(a for a in eng['allowances'] if a['code'] == 'transport')
    assert tr['full_amount'] == 26 and tr['amount'] == pytest.approx(26 - pre)
    assert tr['prorated_days'] == pre
    r = _sheet(conn)[0]['2']
    assert (r['c_trans'], r['c_house']) == (26, 52), 'بالعقد: كاملًا'
    assert r['e_trans'] == pytest.approx(26 - pre) and r['e_house'] == pytest.approx(52 - 2 * pre)
    assert r['e_total'] == pytest.approx(r['e_basic'] + r['e_allow'])
    assert r['net'] == pytest.approx(eng['net'], abs=1e-6)


def test_end_of_service_prorates_allowances_too(env):
    conn, pe = env
    start, end = pe.resolve_period(conn, MONTH, YEAR)
    conn.execute("UPDATE employees SET end_of_service_date = ? WHERE id = 1",
                 ((start + timedelta(days=14)).isoformat(),))
    conn.commit()
    eng = _engine(pe, conn, 1)
    post = next(d for d in eng['deductions'] if d['code'] == 'post_service_days')['days']
    ho = next(a for a in eng['allowances'] if a['code'] == 'housing')
    assert ho['amount'] == pytest.approx(50 - 50 / 26 * post, abs=0.01)
    assert _sheet(conn)[0]['1']['net'] == pytest.approx(eng['net'], abs=1e-6)


def test_a_full_month_keeps_allowances_whole(env):
    conn, pe = env
    for a in _engine(pe, conn, 1)['allowances']:
        if a.get('source') == 'fixed':
            assert 'full_amount' not in a


def test_the_thirty_day_basis_prorates_by_calendar_days(env):
    conn, pe = env
    conn.execute("UPDATE salary_settings_v2 SET setting_value = '30' WHERE setting_name = 'daily_rate_basis'")
    if not conn.execute("SELECT changes()").fetchone()[0]:
        conn.execute("INSERT INTO salary_settings_v2 (setting_name, setting_value) VALUES ('daily_rate_basis', '30')")
    conn.commit()
    _add_fixed(conn, 2, 'transport', 30)
    eng = _engine(pe, conn, 2)
    pre = next(d for d in eng['deductions'] if d['code'] == 'pre_hire_days')['days']
    start, _ = pe.resolve_period(conn, MONTH, YEAR)
    assert pre == (date.fromisoformat('2026-08-10') - start).days
    tr = next(a for a in eng['allowances'] if a['code'] == 'transport')
    assert tr['amount'] == pytest.approx(30 - pre)
    assert _sheet(conn)[0]['2']['work_days'] == 30 - pre
