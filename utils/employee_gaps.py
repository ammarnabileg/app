# -*- coding: utf-8 -*-
"""استكمالُ بيانات الموظّفين الناقصة — ملفُّ Excel يُنزَّل ويُملأ ويُرفع.

ما يُعلَّم بالأحمر في كشف الرواتب (`payroll_sheet.missing_cells`) يُستكمل هنا
دفعةً واحدة بدل فتح ملفّ كلّ موظّف:

1. **التنزيل:** صفٌّ لكلّ موظّفٍ نشطٍ ناقص البيانات، بقيمه الحاليّة، والخانةُ
   الناقصة حمراء. تاريخُ تعيينٍ يشترك فيه كثيرون (غالبًا تاريخُ الاستيراد) أصفر.
2. **الرفع:** تُطبَّق كلُّ خانةٍ غيرُ فارغة تغيّرت، بعد التحقّق من قيمتها،
   ويُسجَّل كلُّ تغييرٍ في سجلّ تغييرات الموظّف.

## الراتب

يُملأ هنا **الناقصُ فقط** (راتبٌ صفر): يصير هو الراتبَ من البداية، في الملفّ
وفي تاريخ الراتب. أمّا تغييرُ راتبٍ قائم فزيادةٌ لها تاريخ سريان، وتكون من
شاشة الموظّف — فلا يعدّل ملفٌّ جماعيٌّ رواتبَ شهورٍ مضت بلا قصد.

## طريقة الدفع

`payment_method`: تحويلٌ بنكيّ، أو نقدًا، أو شيك. تحويلٌ بنكيّ بلا رقم حساب
أو IBAN ناقصٌ، وكذلك موظّفٌ بلا طريقة دفعٍ ولا حساب.
"""
import io
from datetime import date, datetime

PAYMENT_METHODS = {'bank': 'Bank Transfer', 'cash': 'Cash', 'check': 'Check'}
PAYMENT_BY_LABEL = {v.lower(): k for k, v in PAYMENT_METHODS.items()}
PAYMENT_BY_LABEL.update({'تحويل بنكي': 'bank', 'نقدا': 'cash', 'نقدًا': 'cash', 'نقدي': 'cash',
                         'شيك': 'check', 'bank': 'bank', 'cash': 'cash', 'check': 'check',
                         'cheque': 'check'})
CONTRACT_TYPES = ('full_time', 'part_time', 'fixed_term', 'open_ended', 'freelance')
SHARED_HIRE_MIN = 4          # تاريخُ تعيينٍ يتكرّر لهذا العدد فأكثر يُراجَع

# (المفتاح، العمود، عنوان الملف، النوع)
FIELDS = (
    ('arabic_name', 'arabic_name', 'الاسم العربي', 'text'),
    ('hire_date', 'hire_date', 'تاريخ التعيين', 'date'),
    ('department', 'department', 'القسم', 'text'),
    ('position', 'position', 'الوظيفة', 'text'),
    ('cost_center', 'cost_center', 'مركز التكلفة', 'text'),
    ('branch_location', 'branch_location', 'الموقع', 'text'),
    ('contract_type', 'contract_type', 'نوع العقد', 'contract'),
    ('salary', 'salary', 'الراتب الأساسي', 'salary'),
    ('payment_method', 'payment_method', 'طريقة الدفع', 'payment'),
    ('bank_name', 'bank_name', 'البنك', 'text'),
    ('bank_account_number', 'bank_account_number', 'رقم الحساب', 'text'),
    ('bank_iban', 'bank_iban', 'IBAN', 'text'),
)
HEAD_NUM, HEAD_NAME = 'رقم الملف', 'الاسم'
RED, YELLOW, GREY = 'FFFF8B8B', 'FFFFE699', 'FFE7E6E6'


def _get(emp, key):
    try:
        return emp[key] if key in emp.keys() else None
    except AttributeError:
        return emp.get(key)


def has_account(emp):
    return bool(str(_get(emp, 'bank_account_number') or '').strip()
                or str(_get(emp, 'bank_iban') or '').strip())


def payment_label(emp):
    """طريقةُ الدفع كما تُكتب في الكشف، ومعرفةُ النظام بها: (النصّ، معروفة؟)."""
    method = (_get(emp, 'payment_method') or '').strip()
    if method in PAYMENT_METHODS:
        return PAYMENT_METHODS[method], method != 'bank' or has_account(emp)
    return ('Bank Transfer', True) if has_account(emp) else ('Cash', False)


def shared_hire_dates(conn):
    return {r[0] for r in conn.execute(
        'SELECT hire_date FROM employees WHERE is_active = 1 AND hire_date IS NOT NULL '
        "AND hire_date != '' GROUP BY hire_date HAVING COUNT(*) >= ?", (SHARED_HIRE_MIN,))}


def gaps(emp, shared=()):
    """{المفتاح: 'missing' | 'review'} لموظّف."""
    out = {}
    for key, col, _h, kind in FIELDS:
        if kind in ('payment', 'salary') or key.startswith('bank_'):
            continue
        if not str(_get(emp, col) or '').strip():
            out[key] = 'missing'
    if not float(_get(emp, 'salary') or 0) > 0:
        out['salary'] = 'missing'
    _label, known = payment_label(emp)
    if not known:
        out['payment_method'] = 'missing'
        method = (_get(emp, 'payment_method') or '').strip()
        if method == 'bank' or not method:
            out['bank_account_number'] = 'missing'
    hd = str(_get(emp, 'hire_date') or '')[:10]
    if hd and hd in shared:
        out['hire_date'] = 'review'
    return out


def employees_with_gaps(conn):
    shared = shared_hire_dates(conn)
    rows = conn.execute('SELECT * FROM employees WHERE is_active = 1 '
                        'ORDER BY CAST(employee_number AS INTEGER), employee_number').fetchall()
    return [(e, g) for e in rows for g in [gaps(e, shared)] if g]


def summary(conn):
    """{المفتاح: عدد الموظّفين} للشاشة."""
    out = {}
    for _e, g in employees_with_gaps(conn):
        for k, state in g.items():
            if state == 'missing':
                out[k] = out.get(k, 0) + 1
            else:
                out['review_' + k] = out.get('review_' + k, 0) + 1
    return out


# ------------------------------------------------------------------ التنزيل

def export_workbook(conn):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = 'البيانات الناقصة'
    ws.sheet_view.rightToLeft = True
    head = [HEAD_NUM, HEAD_NAME] + [h for _k, _c, h, _t in FIELDS]
    ws.append(head)
    thin = Side(style='thin', color='BBBBBB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i in range(1, len(head) + 1):
        c = ws.cell(1, i)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F5C4A')
        c.alignment, c.border = Alignment(horizontal='center', vertical='center', wrap_text=True), border
    red, yellow, grey = (PatternFill('solid', fgColor=x) for x in (RED, YELLOW, GREY))
    for emp, g in employees_with_gaps(conn):
        vals = [emp['employee_number'], emp['name'] or emp['arabic_name']]
        for key, col, _h, kind in FIELDS:
            v = _get(emp, col)
            if kind == 'payment':
                v = PAYMENT_METHODS.get((v or '').strip()) or None
            elif kind == 'salary':
                v = float(v or 0) or None
            elif kind == 'date' and v:
                try:
                    v = datetime.strptime(str(v)[:10], '%Y-%m-%d')
                except ValueError:
                    pass
            vals.append(v if v not in ('',) else None)
        ws.append(vals)
        r = ws.max_row
        for i in range(1, len(head) + 1):
            ws.cell(r, i).border = border
        ws.cell(r, 1).fill = ws.cell(r, 2).fill = grey
        for j, (key, _c, _h, kind) in enumerate(FIELDS, start=3):
            cell = ws.cell(r, j)
            if kind == 'date':
                cell.number_format = 'yyyy-mm-dd'
            state = g.get(key)
            if state == 'missing':
                cell.fill = red
            elif state == 'review':
                cell.fill = yellow
                cell.comment = Comment('تاريخ تعيين مشترك بين عدة موظفين — غالبًا تاريخ الاستيراد. '
                                       'صحّحه إن كان غير صحيح.', 'onz.one')
    last = max(ws.max_row, 2)
    col_of = {key: j for j, (key, *_r) in enumerate(FIELDS, start=3)}
    from openpyxl.utils import get_column_letter as L
    dv_pay = DataValidation(type='list', formula1='"Bank Transfer,Cash,Check"', allow_blank=True)
    dv_con = DataValidation(type='list', formula1='"' + ','.join(CONTRACT_TYPES) + '"', allow_blank=True)
    ws.add_data_validation(dv_pay)
    ws.add_data_validation(dv_con)
    dv_pay.add(f'{L(col_of["payment_method"])}2:{L(col_of["payment_method"])}{last + 200}')
    dv_con.add(f'{L(col_of["contract_type"])}2:{L(col_of["contract_type"])}{last + 200}')
    widths = [10, 28] + [16] * len(FIELDS)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    ws.column_dimensions[L(col_of['bank_iban'])].width = 32
    ws.freeze_panes = 'C2'
    ws.row_dimensions[1].height = 32

    info = wb.create_sheet('طريقة الاستخدام')
    info.sheet_view.rightToLeft = True
    for line in (
        'استكمال بيانات الموظفين — onz.one',
        '',
        'الأحمر: خانة غير مسجلة في النظام. الأصفر: قيمة تحتاج مراجعة (تاريخ تعيين مشترك بين عدة موظفين).',
        'اكتب القيمة في خانتها ثم ارفع الملف من شاشة «استكمال البيانات». الخانة الفارغة لا تغيّر شيئًا.',
        'لا تغيّر عمود «رقم الملف»: به يُعرف الموظف.',
        'تاريخ التعيين بصيغة 2024-01-31.',
        'طريقة الدفع: Bank Transfer أو Cash أو Check. التحويل البنكي يحتاج رقم الحساب أو IBAN.',
        'نوع العقد: ' + '، '.join(CONTRACT_TYPES) + '.',
        'الراتب الأساسي يُملأ هنا للموظف الذي راتبه صفر فقط. تغيير راتب قائم يكون من شاشة الموظف بتاريخ سريان.',
        'كل تغيير يُسجَّل في سجل تغييرات الموظف.',
    ):
        info.append([line])
    info['A1'].font = Font(bold=True, size=14)
    info.column_dimensions['A'].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ------------------------------------------------------------------ الرفع

def _clean(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _parse(kind, raw):
    """(القيمة، خطأ)."""
    if kind == 'date':
        if isinstance(raw, datetime):
            d = raw.date()
        elif isinstance(raw, date):
            d = raw
        else:
            s = _clean(raw)[:10]
            d = None
            for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
                try:
                    d = datetime.strptime(s, fmt).date()
                    break
                except ValueError:
                    continue
            if d is None:
                return None, 'تاريخ غير صالح'
        if d > date.today():
            return None, 'تاريخ تعيين في المستقبل'
        if d.year < 1950:
            return None, 'تاريخ غير صالح'
        return d.isoformat(), None
    if kind == 'salary':
        try:
            v = float(raw)
        except (TypeError, ValueError):
            return None, 'رقم غير صالح'
        if v <= 0:
            return None, 'الراتب يجب أن يكون أكبر من صفر'
        return round(v, 3), None
    if kind == 'payment':
        m = PAYMENT_BY_LABEL.get(_clean(raw).lower())
        return (m, None) if m else (None, 'طريقة دفع غير معروفة')
    if kind == 'contract':
        v = _clean(raw)
        return (v, None) if v in CONTRACT_TYPES else (None, 'نوع عقد غير معروف')
    return _clean(raw), None


def import_workbook(conn, fileobj, user_id=None):
    """يطبّق الملف. يعيد {'employees', 'fields', 'errors': [(صف، رقم، الخانة، السبب)], 'skipped': [...]}."""
    from openpyxl import load_workbook
    from utils.payroll_engine import log_employee_changes
    from utils.pay_history import BASE_DATE
    wb = load_workbook(fileobj, data_only=True)
    ws = wb.worksheets[0]
    head = [_clean(c.value) for c in ws[1]]
    if HEAD_NUM not in head:
        raise ValueError('الملف ليس ملف استكمال البيانات: لا يوجد عمود «رقم الملف».')
    idx = {h: i for i, h in enumerate(head)}
    cols = [(key, col, kind, idx[h]) for key, col, h, kind in FIELDS if h in idx]
    res = {'employees': 0, 'fields': 0, 'errors': [], 'skipped': []}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        num = _clean(row[idx[HEAD_NUM]] if idx[HEAD_NUM] < len(row) else None)
        if not num:
            continue
        emp = conn.execute('SELECT * FROM employees WHERE employee_number = ?', (num,)).fetchone()
        if emp is None:
            res['errors'].append((r, num, HEAD_NUM, 'موظف غير موجود'))
            continue
        changes = {}
        for key, col, kind, i in cols:
            raw = row[i] if i < len(row) else None
            if _clean(raw) == '' and not isinstance(raw, (date, datetime)):
                continue
            val, err = _parse(kind, raw)
            if err:
                res['errors'].append((r, num, dict((k, h) for k, _c, h, _t in FIELDS)[key], err))
                continue
            cur = _get(emp, col)
            if kind == 'salary':
                if float(cur or 0) > 0:
                    if abs(float(cur) - val) > 0.0005:
                        res['skipped'].append((r, num, 'الراتب الأساسي',
                                               'الراتب مسجّل — التغيير من شاشة الموظف بتاريخ سريان'))
                    continue
            elif _clean(cur) == _clean(val):
                continue
            changes[col] = val
        if not changes:
            continue
        sets = ', '.join(f'{c} = ?' for c in changes)
        conn.execute(f'UPDATE employees SET {sets} WHERE id = ?', list(changes.values()) + [emp['id']])
        if 'salary' in changes:
            # راتبٌ كان صفرًا: هو الراتبُ من البداية، لا زيادةٌ من اليوم
            conn.execute('UPDATE employee_salary_history SET salary = ? WHERE employee_id = ? '
                         'AND COALESCE(salary, 0) <= 0', (changes['salary'], emp['id']))
            if not conn.execute('SELECT 1 FROM employee_salary_history WHERE employee_id = ?',
                                (emp['id'],)).fetchone():
                conn.execute('INSERT INTO employee_salary_history (employee_id, salary, effective_from, '
                             'created_by) VALUES (?, ?, ?, ?)',
                             (emp['id'], changes['salary'], BASE_DATE, user_id))
        after = conn.execute('SELECT * FROM employees WHERE id = ?', (emp['id'],)).fetchone()
        log_employee_changes(conn, emp['id'], emp, after, user_id=user_id, source='data_gaps',
                             note='استكمال البيانات من ملف Excel')
        res['employees'] += 1
        res['fields'] += len(changes)
    return res
