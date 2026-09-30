# -*- coding: utf-8 -*-
"""قسيمةُ الراتب، وملفُّ تحويل البنك، ومخصّصاتُ نهاية الخدمة والإجازات.

## القسيمة

- **للإدارة:** من الكشف المحفوظ إن حُفظ الشهر — ما صُرف فعلًا — وإلّا من
  المحرّك حيًّا بعلامة «مسودّة».
- **للموظّف في البوّابة:** من الكشف **المحفوظ** وحده — ما اعتُمد وحُفظ، لا
  حسابٌ حيّ قد يتغيّر بعد بصمةٍ تُصحَّح. وقبل الحفظ لا قسيمة.

## ملفُّ البنك

صافي كلّ موظّف له حساب، بمرجع «Salary MM/YYYY»؛ ومن لا حساب له في ورقةٍ
ثانية (نقدًا). من الكشف المحفوظ إن وُجد، وإلّا من الحساب الحيّ ويُكتب ذلك
في رأس الملفّ. صيغةٌ عامّة تُرفع أو تُحوَّل لصيغة كلّ بنك.

## المخصّصات

ما على الشركة لو انتهت خدمةُ كلّ موظّف في تاريخٍ ما: مكافأةُ نهاية الخدمة
(إنهاءٌ من صاحب العمل — أعلى الالتزام) ورصيدُ الإجازات بأجر اليوم. من دالّة
نهاية الخدمة نفسِها (`compute_kuwait_eos`) فلا يختلف المخصّصُ عن التسوية.
"""
import io
import json
from datetime import date


def _emp_rows(conn, ids):
    if not ids:
        return {}
    q = 'SELECT * FROM employees WHERE id IN (%s)' % ','.join('?' * len(ids))
    return {r['id']: r for r in conn.execute(q, list(ids)).fetchall()}


def _leave_balance(conn, emp_id, as_of):
    try:
        from utils.leave_balance import compute_leave_balance
        b = compute_leave_balance(conn, emp_id, as_of=as_of)
        return b['balance'] if b else None
    except Exception:
        return None


def saved_run(conn, month, year):
    return conn.execute('SELECT * FROM payroll_runs WHERE month = ? AND year = ?',
                        (int(month), int(year))).fetchone()


def slips_for(conn, month, year, employee_id=None):
    """المحفوظُ إن وُجد كشفٌ للشهر، وإلّا الحيّ."""
    if saved_run(conn, month, year):
        return slips_saved(conn, month, year, employee_id)
    return slips_live(conn, month, year, employee_id)


def slips_live(conn, month, year, employee_id=None):
    from utils.payroll_engine import compute_monthly_payroll
    data = compute_monthly_payroll(conn, month, year)
    rows = [r for r in data['rows'] if employee_id is None or r['employee_id'] == int(employee_id)]
    emps = _emp_rows(conn, [r['employee_id'] for r in rows])
    out = []
    for r in rows:
        out.append(_slip(conn, emps.get(r['employee_id']), data['period'], r['basic'],
                         r['allowances'], r['deductions'], r['total_allowances'],
                         r['total_deductions'], r['net'], r['required_hours'],
                         r['actual_hours'], saved=False))
    return out


def slips_saved(conn, month, year, employee_id=None):
    run = saved_run(conn, month, year)
    if not run:
        return []
    q = 'SELECT * FROM payroll_run_lines WHERE run_id = ?'
    args = [run['id']]
    if employee_id is not None:
        q += ' AND employee_id = ?'
        args.append(int(employee_id))
    lines = conn.execute(q, args).fetchall()
    emps = _emp_rows(conn, [ln['employee_id'] for ln in lines])
    period = {'start': run['period_start'], 'end': run['period_end']}
    out = []
    for ln in lines:
        det = json.loads(ln['details_json'] or '{}')
        out.append(_slip(conn, emps.get(ln['employee_id']), period, ln['basic'],
                         det.get('allowances', []), det.get('deductions', []),
                         ln['total_allowances'], ln['total_deductions'], ln['net'],
                         ln['required_hours'], ln['actual_hours'], saved=True,
                         status=run['status']))
    return out


def _slip(conn, emp, period, basic, allowances, deductions, t_allow, t_ded, net,
          req_h, act_h, saved, status=None):
    e = emp
    ek = e.keys() if e is not None else []
    return {
        'employee_id': e['id'] if e is not None else None,
        'employee_number': e['employee_number'] if e is not None else '',
        'name': (e['arabic_name'] or e['name']) if e is not None else '',
        'name_en': e['name'] if e is not None else '',
        'department': (e['department'] if e is not None else '') or '',
        'position': (e['position'] if e is not None else '') or '',
        'hire_date': (e['hire_date'] if e is not None else '') or '',
        'bank': (e['bank_name'] if e is not None and 'bank_name' in ek else '') or '',
        'account': (e['bank_account_number'] if e is not None and 'bank_account_number' in ek else '') or '',
        'period': period, 'basic': basic,
        'allowances': allowances, 'deductions': deductions,
        'total_allowances': t_allow, 'total_deductions': t_ded, 'net': net,
        'required_hours': req_h, 'actual_hours': act_h,
        'leave_balance': _leave_balance(conn, e['id'], period['end']) if e is not None else None,
        'saved': saved, 'status': status,
    }


# ------------------------------------------------------------------ البنك

def bank_workbook(conn, month, year, company=''):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    run = saved_run(conn, month, year)
    slips = slips_saved(conn, month, year) if run else slips_live(conn, month, year)
    source = 'saved' if run else 'live'
    ref = f'Salary {month:02d}/{year}'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Bank'
    cash = wb.create_sheet('Cash')
    head_fill = PatternFill('solid', fgColor='1F5C4A')
    head = ['Emp#', 'Name', 'Civil ID', 'Bank', 'Account / IBAN', 'Amount', 'Reference']
    for sheet, title in ((ws, 'Bank transfer'), (cash, 'Cash / no account')):
        sheet.append([f'{company} — {title} — {ref}'
                      + ('' if source == 'saved' else '  (NOT SAVED: live figures)')])
        sheet.cell(1, 1).font = Font(bold=True, size=12,
                                     color='000000' if source == 'saved' else 'C00000')
        sheet.append([])
        sheet.append(head)
        for i in range(1, len(head) + 1):
            c = sheet.cell(3, i)
            c.font, c.fill = Font(bold=True, color='FFFFFF'), head_fill
            c.alignment = Alignment(horizontal='center')
    emps = _emp_rows(conn, [s['employee_id'] for s in slips])
    tot = {'Bank': 0.0, 'Cash': 0.0}
    for s in slips:
        if (s['net'] or 0) <= 0.0005:
            continue
        e = emps.get(s['employee_id'])
        civil = (e['national_id'] if e is not None and 'national_id' in e.keys() else '') or ''
        row = [s['employee_number'], s['name_en'] or s['name'], civil, s['bank'], s['account'],
               round(float(s['net']), 3), ref]
        target = ws if s['account'].strip() else cash
        target.append(row)
        target.cell(target.max_row, 6).number_format = '#,##0.000'
        tot[target.title] += float(s['net'])
    for sheet in (ws, cash):
        sheet.append([])
        sheet.append(['', 'Total', '', '', '', round(tot[sheet.title], 3), ''])
        sheet.cell(sheet.max_row, 6).number_format = '#,##0.000'
        sheet.cell(sheet.max_row, 2).font = Font(bold=True)
        for col, w in zip('ABCDEFG', (8, 30, 14, 16, 30, 13, 16)):
            sheet.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, source


# ------------------------------------------------------------------ المخصّصات

def provisions(conn, as_of):
    from routes.eos_routes import compute_kuwait_eos
    rows = []
    for e in conn.execute("SELECT id, employee_number, name, arabic_name, department, hire_date "
                          "FROM employees WHERE is_active = 1 "
                          "ORDER BY CAST(employee_number AS INTEGER), employee_number").fetchall():
        c = compute_kuwait_eos(conn, e['id'], as_of, 'termination')
        if not c:
            continue
        rows.append({'emp_no': e['employee_number'], 'name': e['arabic_name'] or e['name'],
                     'department': e['department'] or '', 'hire_date': e['hire_date'],
                     'years': c['years'], 'monthly_wage': c['monthly_wage'],
                     'daily_rate': c['daily_rate'], 'eos': c['gross_gratuity'],
                     'leave_days': c['leave_days'], 'leave_amount': c['leave_amount'],
                     'total': round(c['gross_gratuity'] + c['leave_amount'], 3)})
    return rows


def provisions_workbook(conn, as_of, company=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = provisions(conn, as_of)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Provisions'
    ws.append([f'{company} — End of service & leave provisions as of {as_of}'])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append(['Termination by the employer (Kuwait Labour Law 6/2010, Art. 51) — '
               'the highest liability; resignation reduces it.'])
    ws.append([])
    head = ['Emp#', 'Name', 'Dept.', 'DOJ', 'Years', 'Monthly wage', 'Daily rate',
            'EOS provision', 'Leave days', 'Leave provision', 'Total']
    ws.append(head)
    for i in range(1, len(head) + 1):
        c = ws.cell(4, i)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F5C4A')
    keys = ['emp_no', 'name', 'department', 'hire_date', 'years', 'monthly_wage', 'daily_rate',
            'eos', 'leave_days', 'leave_amount', 'total']
    for r in rows:
        ws.append([r[k] for k in keys])
        for i in (6, 7, 8, 10, 11):
            ws.cell(ws.max_row, i).number_format = '#,##0.000'
        ws.cell(ws.max_row, 5).number_format = '0.00'
    ws.append([])
    ws.append(['', 'Total', '', '', '', round(sum(r['monthly_wage'] for r in rows), 3), '',
               round(sum(r['eos'] for r in rows), 3), round(sum(r['leave_days'] for r in rows), 2),
               round(sum(r['leave_amount'] for r in rows), 3), round(sum(r['total'] for r in rows), 3)])
    for i in (2, 6, 8, 10, 11):
        ws.cell(ws.max_row, i).font = Font(bold=True)
    for col, w in zip('ABCDEFGHIJK', (8, 28, 16, 12, 8, 13, 11, 14, 10, 14, 14)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, rows


def today_iso():
    return date.today().isoformat()
