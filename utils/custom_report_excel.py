# -*- coding: utf-8 -*-
"""Excel التقرير الشامل المخصّص (/fingerprint/custom_employees_report) — منسّق من السيرفر.

كان التصدير من المتصفح (SheetJS) بينسخ جدول الصفحة زي ما هو: من غير ألوان ولا عرض أعمدة، والساعات
نص «8:30» مايتجمعش، وأيام الشهر بتطلع عمود واحد مدموج. هنا الملف بيتبني من نفس بيانات الصفحة
(نفس الـ API) بـ openpyxl:

- **ورقة «التقرير الشامل»:** عنوان واسم الشركة والفترة، وصف عناوين من سطرين (التاريخ واليوم)،
  وكل يوم برمزه ولونه زي الصفحة، والساعات **وقت حقيقي** ([h]:mm) فبتتجمع، وسطر إجمالي بمعادلات،
  وتجميد للعناوين والاسم، وطباعة A4 بالعرض على عرض صفحة واحدة، والعناوين بتتكرر في كل صفحة.
- **ورقة «تفاصيل الحضور»:** سطر لكل موظف ويوم: الحالة بالعربي والدخول والخروج والساعات، بفلتر.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = 'Calibri'
C_TITLE = '1F3B57'
C_HEAD = '2F4F6F'
C_HEAD_DAYS = '3E6A8A'
C_HOURS = '4A5568'
C_MONEY = '1F6F5C'
C_DED = '9B2C2C'
C_NET = '1E4E8C'
C_ZEBRA = 'F5F8FB'
C_TOTAL = 'E3ECF5'

STATUS = {   # الحالة ← (الرمز، لون الخلفية، لون الخط، الاسم بالعربي)
    'holiday':    ('إ', 'FFF1C2', '8A6100', 'عطلة رسمية'),
    'leave':      ('ج', 'D6EEF8', '1F6F8B', 'إجازة'),
    'weekly_off': ('ر', 'ECEEF1', '6B7280', 'راحة أسبوعية'),
    'absent':     ('غ', 'FBD5D5', 'B42318', 'غياب'),
    'partial':    ('ن', 'D9534F', 'FFFFFF', 'بصمة ناقصة'),
    'terminated': ('×', 'E5E7EB', '6B7280', 'انتهت الخدمة'),
    'not_hired':  ('–', 'F9FAFB', 'C0C4CC', 'قبل التعيين'),
    'present':    ('✓', 'FFFFFF', '1A7F37', 'حاضر'),
}

_thin = Side(style='thin', color='C9D3DD')
_med = Side(style='medium', color='7A8CA0')
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True, indent=1)
HM = '[h]:mm'
MONEY = '#,##0.000'
AR_DAYS = ['الاثنين', 'الثلاثاء', 'الأربعاء', 'الخميس', 'الجمعة', 'السبت', 'الأحد']


def _fill(c):
    return PatternFill('solid', fgColor=c)


def _days(h):
    """ساعات ← قيمة وقت في Excel (كسر من يوم)، فتتجمع وتتعرض [h]:mm."""
    try:
        return round(float(h or 0) / 24.0, 8)
    except (TypeError, ValueError):
        return 0


def _head(ws, r, c, text, color, r2=None, c2=None, size=10):
    cell = ws.cell(r, c, text)
    cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=size)
    cell.fill = _fill(color)
    cell.alignment = CENTER
    cell.border = BORDER
    if r2 or c2:
        ws.merge_cells(start_row=r, start_column=c, end_row=r2 or r, end_column=c2 or c)
        for rr in range(r, (r2 or r) + 1):
            for cc in range(c, (c2 or c) + 1):
                x = ws.cell(rr, cc)
                x.fill, x.border = _fill(color), BORDER
    return cell


def build(data, company='', dept='', pos='', labels=None):
    """data: ناتج /api/fingerprint/custom_employees_report. labels: عناوين مترجمة اختيارية."""
    L = {'title': 'التقرير الشامل للموظفين', 'num': 'رقم الملف', 'name': 'الاسم', 'dept': 'القسم',
         'pos': 'المسمى الوظيفي', 'work_days': 'أيام العمل', 'sick': 'مرضي / طارئ',
         'req': 'الساعات المطلوبة', 'act': 'الساعات الفعلية', 'miss': 'الساعات الناقصة',
         'ot': 'ساعات الإضافي', 'total_h': 'إجمالي الساعات', 'late': 'التأخير',
         'present': 'أيام الحضور', 'absent': 'أيام الغياب', 'basic': 'الراتب الأساسي',
         'allow': 'البدلات', 'ot_val': 'قيمة الإضافي', 'ded': 'إجمالي الخصومات', 'net': 'صافي الراتب',
         'details': 'تفاصيل الحضور', 'date': 'التاريخ', 'day': 'اليوم', 'status': 'الحالة',
         'in': 'الدخول', 'out': 'الخروج', 'hours': 'ساعات العمل', 'total': 'الإجمالي'}
    L.update(labels or {})
    rows = data.get('rows') or []
    start, end = data.get('start_date', ''), data.get('end_date', '')
    dates = [c['date'] for c in rows[0]['daily_cells']] if rows else []

    wb = Workbook()
    ws = wb.active
    ws.title = L['title'][:31]
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

    info = [('num', 9), ('name', 28), ('dept', 16), ('pos', 18), ('work_days', 8), ('sick', 8)]
    hours = [('req', 'required_hours'), ('act', 'basic_hours'), ('miss', 'missing_hours'),
             ('ot', 'overtime_hours'), ('total_h', 'total_work_hours'), ('late', None),
             ('present', 'present_days'), ('absent', 'absent_days')]
    money = [('basic', C_MONEY), ('allow', C_MONEY), ('ot_val', C_MONEY), ('ded', C_DED), ('net', C_NET)]
    n_info, n_days = len(info), len(dates)
    first_day = n_info + 1
    first_hours = first_day + n_days
    first_money = first_hours + len(hours)
    last_col = first_money + len(money) - 1

    # --- العنوان
    ws.row_dimensions[1].height = 30
    t = ws.cell(1, 1, company or L['title'])
    t.font = Font(name=FONT, bold=True, size=16, color=C_TITLE)
    t.alignment = Alignment(horizontal='right', vertical='center')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(last_col, 12))
    sub = f"{L['title']}  —  من {start} إلى {end}"
    if dept:
        sub += f"  —  {L['dept']}: {dept}"
    if pos:
        sub += f"  —  {L['pos']}: {pos}"
    s = ws.cell(2, 1, sub)
    s.font = Font(name=FONT, size=11, color='4B5563')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(last_col, 12))
    # الدليل
    col = 1
    lg = ws.cell(3, col, 'دليل الرموز:')
    lg.font = Font(name=FONT, bold=True, size=9, color='4B5563')
    widths = {i: w for i, (_, w) in enumerate(info, start=1)}
    for key in ('present', 'weekly_off', 'holiday', 'leave', 'absent', 'partial', 'terminated'):
        sym, bg, fg, name = STATUS[key]
        col += 1
        span, w = col, widths.get(col, 6.2)
        while w < 12 and span < last_col:           # الأعمدة الضيقة (أيام الشهر): البند ياخد كذا عمود
            span += 1
            w += widths.get(span, 6.2)
        if span > last_col:
            break
        c = ws.cell(3, col, f'{sym} {name}')
        c.font = Font(name=FONT, size=9, color=fg, bold=True)
        for cc in range(col, span + 1):
            x = ws.cell(3, cc)
            x.fill, x.border = _fill(bg), BORDER
        c.alignment = CENTER
        if span > col:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=span)
        col = span
    ws.row_dimensions[3].height = 26

    # --- العناوين (سطرين: 5 و6)
    H1, H2 = 5, 6
    ws.row_dimensions[H1].height = 22
    ws.row_dimensions[H2].height = 30
    for i, (k, w) in enumerate(info, start=1):
        _head(ws, H1, i, L[k], C_HEAD, r2=H2)
        ws.column_dimensions[get_column_letter(i)].width = w
    for j, d in enumerate(dates):
        c = first_day + j
        try:
            dt = datetime.strptime(d, '%Y-%m-%d')
            top, low = dt.strftime('%d/%m'), AR_DAYS[dt.weekday()]
            wkend = dt.weekday() in (4, 5)
        except ValueError:
            top, low, wkend = d, '', False
        _head(ws, H1, c, top, '5B6B7C' if wkend else C_HEAD_DAYS, size=9)
        _head(ws, H2, c, low, '5B6B7C' if wkend else C_HEAD_DAYS, size=8)
        ws.column_dimensions[get_column_letter(c)].width = 6.2
    for j, (k, _) in enumerate(hours):
        c = first_hours + j
        _head(ws, H1, c, L[k], C_DED if k in ('miss', 'absent', 'late') else C_HOURS, r2=H2, size=9)
        ws.column_dimensions[get_column_letter(c)].width = 10
    for j, (k, color) in enumerate(money):
        c = first_money + j
        _head(ws, H1, c, L[k], color, r2=H2)
        ws.column_dimensions[get_column_letter(c)].width = 13

    # --- الصفوف
    r = H2
    for n, row in enumerate(rows):
        r += 1
        ws.row_dimensions[r].height = 30 if row.get('arabic_name') and row['arabic_name'] != row['name'] else 20
        zebra = _fill(C_ZEBRA) if n % 2 else None
        st, fi = row['stats'], row['financials']
        name = row['name'] or ''
        if row.get('arabic_name') and row['arabic_name'] != name:
            name = f"{row['arabic_name']}\n{name}"
        num = row.get('employee_number') or ''
        vals = [int(num) if str(num).isdigit() else num, name, row.get('department') or '',
                row.get('position') or '', st['work_days'], st['sick_days']]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, size=10, bold=i in (1, 2))
            c.alignment = RIGHT if i in (2, 3, 4) else CENTER
            c.border = BORDER
            if zebra:
                c.fill = zebra
        for j, cell in enumerate(row['daily_cells']):
            c = ws.cell(r, first_day + j)
            status = cell.get('status')
            sym, bg, fg, _ = STATUS.get(status, ('', 'FFFFFF', '9CA3AF', ''))
            ot = float(cell.get('overtime_minutes') or 0)
            if status == 'present' and float(cell.get('work_hours') or 0) > 0:
                c.value = _days(cell['work_hours'])
                c.number_format = HM
                c.font = Font(name=FONT, size=9, color='111827')
                if zebra:
                    c.fill = zebra
            elif status in ('weekly_off', 'holiday') and ot > 0:
                c.value = _days(ot / 60.0)           # شغل يوم راحة/عطلة = إضافي: الساعات بلون الإضافي
                c.number_format = HM
                c.font = Font(name=FONT, size=9, bold=True, color='1A7F37')
                c.fill = _fill(bg)
            else:
                c.value = sym or None
                c.font = Font(name=FONT, size=10, bold=True, color=fg)
                if status in STATUS and status != 'present':
                    c.fill = _fill(bg)
                elif zebra:
                    c.fill = zebra
            c.alignment, c.border = CENTER, BORDER
        hv = [_days(st['required_hours']), _days(st['basic_hours']), _days(st['missing_hours']),
              _days(st['overtime_hours']), _days(st['total_work_hours']), _days(st['total_late_mins'] / 60.0),
              st['present_days'], st['absent_days']]
        for j, v in enumerate(hv):
            c = ws.cell(r, first_hours + j, v)
            k = hours[j][0]
            c.number_format = HM if j < 6 else '0'
            c.font = Font(name=FONT, size=10, bold=k in ('total_h',),
                          color='B42318' if k in ('miss', 'late', 'absent') and v else
                          '1A7F37' if k in ('ot', 'present') and v else '111827')
            c.alignment, c.border = CENTER, BORDER
            if zebra:
                c.fill = zebra
        mv = [row['salary'], fi.get('allowances', 0), fi['overtime_reward'], fi['total_deduction'], fi['net_salary']]
        for j, v in enumerate(mv):
            c = ws.cell(r, first_money + j, round(float(v or 0), 3))
            c.number_format = MONEY
            k = money[j][0]
            c.font = Font(name=FONT, size=10, bold=k == 'net',
                          color={'allow': '1A7F37', 'ot_val': '1A7F37', 'ded': 'B42318', 'net': C_NET}.get(k, '111827'))
            c.alignment, c.border = CENTER, BORDER
            c.fill = _fill('EAF2FB') if k == 'net' else (zebra or PatternFill())
    last_row = r

    # --- الإجمالي
    if rows:
        r += 1
        ws.row_dimensions[r].height = 24
        tc = ws.cell(r, 1, f"{L['total']} ({len(rows)} موظف)")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_info)
        for cc in range(1, last_col + 1):
            x = ws.cell(r, cc)
            x.fill, x.border = _fill(C_TOTAL), Border(left=_thin, right=_thin, top=_med, bottom=_med)
            x.font = Font(name=FONT, bold=True, size=10, color=C_TITLE)
            x.alignment = CENTER
        tc.alignment = RIGHT
        for cc in list(range(first_hours, first_money + len(money))):
            L_ = get_column_letter(cc)
            x = ws.cell(r, cc, f'=SUM({L_}{H2 + 1}:{L_}{last_row})')
            x.number_format = (HM if cc < first_hours + 6 else '0') if cc < first_money else MONEY

    # --- العرض والطباعة
    ws.freeze_panes = ws.cell(H2 + 1, 3)
    ws.print_title_rows = f'{H1}:{H2}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.oddFooter.center.text = 'صفحة &P من &N'
    ws.oddHeader.right.text = (company or '')[:80]

    _details(wb, rows, L, company, start, end)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _details(wb, rows, L, company, start, end):
    ws = wb.create_sheet(L['details'][:31])
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    cols = [(L['num'], 10), (L['name'], 28), (L['dept'], 16), (L['date'], 12), (L['day'], 11),
            (L['status'], 15), (L['in'], 10), (L['out'], 10), (L['hours'], 11)]
    t = ws.cell(1, 1, f"{company + '  —  ' if company else ''}{L['details']}  —  من {start} إلى {end}")
    t.font = Font(name=FONT, bold=True, size=14, color=C_TITLE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.row_dimensions[1].height = 26
    for i, (h, w) in enumerate(cols, start=1):
        _head(ws, 3, i, h, C_HEAD)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 22
    r = 3
    for n, row in enumerate(rows):
        num = row.get('employee_number') or ''
        for cell in row['daily_cells']:
            r += 1
            st = cell.get('status')
            sym, bg, fg, name = STATUS.get(st, ('', 'FFFFFF', '111827', st or ''))
            try:
                dt = datetime.strptime(cell['date'], '%Y-%m-%d')
                day = AR_DAYS[dt.weekday()]
            except (KeyError, ValueError):
                dt, day = cell.get('date'), ''
            cin = cell.get('check_in') if cell.get('check_in') not in (None, '--', '-') else ''
            cout = cell.get('check_out') if cell.get('check_out') not in (None, '--', '-') else ''
            wh = float(cell.get('work_hours') or 0)
            vals = [int(num) if str(num).isdigit() else num, row.get('arabic_name') or row['name'],
                    row.get('department') or '', dt, day, name, cin, cout, _days(wh) if wh else None]
            band = _fill(C_ZEBRA) if n % 2 else None
            for i, v in enumerate(vals, start=1):
                c = ws.cell(r, i, v)
                c.font = Font(name=FONT, size=10)
                c.alignment = RIGHT if i in (2, 3) else CENTER
                c.border = BORDER
                if band:
                    c.fill = band
            ws.cell(r, 4).number_format = 'yyyy-mm-dd'
            ws.cell(r, 9).number_format = HM
            s = ws.cell(r, 6)
            if st in STATUS and st != 'present':
                s.fill = _fill(bg)
            s.font = Font(name=FONT, size=10, bold=True, color=fg)
    ws.freeze_panes = 'A4'
    if r > 3:
        ws.auto_filter.ref = f'A3:{get_column_letter(len(cols))}{r}'
    ws.print_title_rows = '3:3'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = 'صفحة &P من &N'
