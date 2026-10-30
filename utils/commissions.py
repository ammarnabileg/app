# -*- coding: utf-8 -*-
"""العمولات: بتيجي من برا (onz.one / قصة / ملف Excel) وتنزل بند «عمولة مبيعات» في كشف الرواتب.

اتفقنا إن الرواتب هي اللي بتدفع العمولة (اختيار «أ»): النظام اللي بيحسبها بيبعت المبلغ بس، وكشف
الرواتب هو اللي بيصرفه — بنفس الموافقة والقفل والدفتر.

## طريقتين للدخول
- **ملف** (Excel أو CSV) من شاشة «العمولات»: عمود رقم الموظف والمبلغ، واختياري: الشهر، السنة،
  المرجع، السبب.
- **API** للوحة onz.one: `POST /api/payroll/commissions` بمفتاح (`X-Commission-Key` أو
  `Authorization: Bearer …`). المفتاح بيتعمل من الشاشة ويظهر مرة واحدة، والمحفوظ hash بس.

## القواعد
- **المرجع بيمنع التكرار:** نفس المرجع لنفس الموظف مايتسجّلش مرتين (إعادة الإرسال آمنة). من غير مرجع
  بيتعمل مرجع من (الموظف، الشهر، المبلغ، رقم السطر في الملف).
- **شهر مقفول مايتلمسش:** العمولة لشهر اتعتمد كشفه بترجع «مقفول» — مابتتنقلش لشهر تاني في صمت.
- **موظف مش معروف** بيرجع في التقرير، وباقي الملف بيكمل.
"""
import csv
import hashlib
import io
import secrets
from datetime import date

ITEM_CODE = 'commission'
KEY_SETTING = 'commissions_api_key_sha256'

_HEAD = {
    'employee_number': ('employee_number', 'رقم الموظف', 'الرقم الوظيفي', 'رقم الملف', 'emp_no', 'pin'),
    'amount': ('amount', 'المبلغ', 'العمولة', 'commission'),
    'month': ('month', 'الشهر'),
    'year': ('year', 'السنة'),
    'reference': ('reference', 'ref', 'المرجع'),
    'reason': ('reason', 'السبب', 'البيان', 'description', 'note'),
}


# ------------------------------------------------------------ المفتاح

def _sha(key):
    return hashlib.sha256(str(key or '').strip().encode()).hexdigest()


def new_api_key():
    """مفتاح جديد — يتعرض مرة واحدة، والمحفوظ hash بس (المفتاح القديم بيبطل)."""
    from utils.db import set_setting
    key = 'cmk_' + secrets.token_urlsafe(32)
    set_setting(KEY_SETTING, _sha(key))
    return key


def revoke_api_key():
    from utils.db import set_setting
    set_setting(KEY_SETTING, '')


def has_api_key():
    from utils.db import get_setting
    return bool(get_setting(KEY_SETTING, '') or '')


def check_api_key(key):
    from utils.db import get_setting
    want = get_setting(KEY_SETTING, '') or ''
    return bool(want and key and secrets.compare_digest(want, _sha(key)))


# ------------------------------------------------------------ الملف

def _norm(h):
    h = str(h or '').strip().lower()
    for k, names in _HEAD.items():
        if h in (n.lower() for n in names):
            return k
    return None


def parse_file(data, filename):
    """[(رقم السطر، {employee_number, amount, …})] من Excel أو CSV."""
    name = (filename or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode('utf-8-sig', errors='replace')
        rows = list(csv.reader(io.StringIO(text)))
    head_i = next((i for i, r in enumerate(rows[:10])
                   if {'employee_number', 'amount'} <= {_norm(c) for c in r}), None)
    if head_i is None:
        raise ValueError('مالقيتش عمود «رقم الموظف» و«المبلغ» في أول 10 سطور')
    cols = [_norm(c) for c in rows[head_i]]
    out = []
    for n, r in enumerate(rows[head_i + 1:], start=head_i + 2):
        item = {k: r[i] for i, k in enumerate(cols) if k and i < len(r)}
        if not any(str(v or '').strip() for v in item.values()):
            continue
        out.append((n, item))
    return out


# ------------------------------------------------------------ التسجيل

def _item_id(conn):
    r = conn.execute('SELECT id FROM payroll_item_types WHERE code = ?', (ITEM_CODE,)).fetchone()
    if not r:
        raise ValueError('بند «عمولة مبيعات» مش موجود في بنود الرواتب')
    return r[0]


def _num(v):
    s = str(v if v is not None else '').strip().replace(',', '')
    s = s.translate(str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789'))
    return float(s)


def import_items(conn, items, user_id=None, month=None, year=None, source='file'):
    """items: [(رقم السطر، dict)] — يُرجع تقرير. مابيعملش commit لو فيه خطأ في الاتصال (المستدعي)."""
    from utils.payroll_engine import is_period_locked
    item_id = _item_id(conn)
    today = date.today()
    rep = {'added': 0, 'total': 0.0, 'duplicate': 0, 'unknown': [], 'locked': [], 'invalid': []}
    for line, it in items:
        pin = str(it.get('employee_number') or '').strip()
        if pin.endswith('.0'):
            pin = pin[:-2]
        try:
            amt = round(_num(it.get('amount')), 3)
            m = int(_num(it.get('month'))) if str(it.get('month') or '').strip() else int(month or today.month)
            y = int(_num(it.get('year'))) if str(it.get('year') or '').strip() else int(year or today.year)
            if not (1 <= m <= 12 and 2000 <= y <= 2100) or amt <= 0:
                raise ValueError
        except (TypeError, ValueError):
            rep['invalid'].append(line)
            continue
        emp = conn.execute('SELECT id FROM employees WHERE CAST(employee_number AS TEXT) = ?', (pin,)).fetchone()
        if not emp:
            rep['unknown'].append((line, pin))
            continue
        if is_period_locked(conn, m, y):
            rep['locked'].append((line, pin, f'{m}/{y}'))
            continue
        ref = str(it.get('reference') or '').strip()[:120] or f'{source}:{pin}:{y}-{m:02d}:{amt}:{line}'
        ref = 'COMM:' + ref
        if conn.execute("SELECT 1 FROM payroll_transactions WHERE employee_id = ? AND reference = ? AND status = 'active'",
                        (emp[0], ref)).fetchone():
            rep['duplicate'] += 1
            continue
        conn.execute('''INSERT INTO payroll_transactions (employee_id, item_type_id, month, year, amount, reason,
                        reference, entered_by) VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                     (emp[0], item_id, m, y, amt, str(it.get('reason') or 'عمولة').strip()[:300], ref, user_id))
        rep['added'] += 1
        rep['total'] = round(rep['total'] + amt, 3)
    conn.commit()
    return rep


def recent(conn, limit=50):
    return [dict(r) for r in conn.execute(
        '''SELECT t.id, e.employee_number, COALESCE(e.arabic_name, e.name) AS name, t.month, t.year, t.amount,
                  t.reason, t.reference, t.status, t.created_at
           FROM payroll_transactions t JOIN payroll_item_types it ON it.id = t.item_type_id
           JOIN employees e ON e.id = t.employee_id
           WHERE it.code = ? ORDER BY t.id DESC LIMIT ?''', (ITEM_CODE, int(limit))).fetchall()]
