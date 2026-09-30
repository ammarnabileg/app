# -*- coding: utf-8 -*-
"""اعتمادُ الساعات إلى Excel: صفٌّ لكلّ موظّف، وعمودٌ لكلّ يومٍ في الدورة.

الأرقامُ من المحرّك نفسه الذي تعرضه الشاشة وتطبعه الطباعة ويصرفه كشفُ
الرواتب — `compute_employee_days` لليوم، و`compute_month_metrics` للمجاميع،
و`compute_monthly_payroll` للمال. لا حسابَ ثانٍ هنا يفترق عنها يومًا.

## خليّةُ اليوم

- يومٌ اشتُغل (عملٌ، أو راحةٌ وعطلةٌ اشتُغلتا): ساعاتُه، قيمةُ وقتٍ حقيقيّة
  بتنسيق `[h]:mm` — 8:51 لا 8.85، وتُجمع في Excel.
- وغيرُه حرفٌ بلونه، ومفتاحُ الحروف تحت الجدول.

## الأعمدةُ بعد الأيّام

- **الفعليّةُ** من أوّل بصمةٍ لآخرها، فتشمل وقتَ الإضافيّ. والإضافيُّ لا يسدّ
  نقصًا، فالعاديّةُ = الفعليّة − الإضافيّ.
- **الناقصة** = المطلوبة − العاديّة، إن كانت موجبة.
- **الإجمالي** = المطلوبة − الناقصة + الإضافيّ: ما يُحتسب للموظّف — بلا عدٍّ
  للإضافيّ مرّتين.
- **المال** (الأساسيّ، قيمةُ الإضافيّ، الخصومات، الصافي) لمن يملك حسابَ
  الرواتب وحده؛ ومن يعتمد الساعاتِ بلا صلاحية الرواتب يأخذ الملفَّ بلا مال.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.timefmt import EXCEL_HOURS_FORMAT, excel_hours

# الحالة ← (الحرف، لون الخلفيّة، مفتاح الترجمة للمفتاح أسفل الجدول)
DAY_CODES = {
    'absent': ('غ', 'F8D7DA', 'x.hx_absent'),
    'partial': ('ن', 'FFE5B4', 'x.hx_partial'),
    'weekly_off': ('ر', 'E9ECEF', 'x.hx_rest'),
    'holiday': ('ع', 'D6EAF8', 'x.hx_holiday'),
    'leave_paid': ('إ', 'FFF3CD', 'x.hx_leave'),
    'leave_sick': ('م', 'FCE4D6', 'x.hx_sick'),
    'leave_unpaid': ('ب', 'F4CCCC', 'x.hx_unpaid'),
    'excused': ('ذ', 'E2F0D9', 'x.hx_excused'),
    'not_hired': ('-', 'FFFFFF', 'x.hx_outside'),
    'ended': ('-', 'FFFFFF', 'x.hx_outside'),
}
WORKED_OFF_FILL = 'C6EFCE'      # ساعاتٌ في يوم راحةٍ أو عطلة
MONEY_FORMAT = '#,##0.000'


def _emergency_names(conn):
    from utils.labor_law import leave_kind
    try:
        return {r['name'] for r in conn.execute('SELECT * FROM leave_types').fetchall()
                if leave_kind(r) == 'emergency'}
    except Exception:
        return set()


def build_hours_workbook(conn, month, year, dept=None, with_money=False, gettext=None):
    from utils.payroll_engine import (compute_employee_days, compute_month_metrics,
                                      compute_monthly_payroll, fetch_payroll_employees,
                                      resolve_period)
    _ = gettext or (lambda s: s)
    p_start, p_end = resolve_period(conn, month, year)
    employees = fetch_payroll_employees(conn, dept=dept or None, period=(p_start, p_end))
    metrics = compute_month_metrics(conn, month, year, employees)
    money = {}
    if with_money:
        money = {r['employee_id']: r
                 for r in compute_monthly_payroll(conn, month, year)['rows']}
    emergency = _emergency_names(conn)

    days = []
    d = p_start
    while d <= p_end:
        days.append(d)
        d = date.fromordinal(d.toordinal() + 1)
    day_names = [_('day.sun'), _('day.mon'), _('day.tue'), _('day.wed'),
                 _('day.thu'), _('day.fri'), _('day.sat')]

    head = [_('x.hx_file_no'), _('x.hx_name'), _('x.hx_department'), _('x.hx_position'),
            _('x.hx_work_days'), _('x.hx_sick_emergency')]
    head += [f"{x.strftime('%d/%m')}{day_names[x.isoweekday() % 7]}" for x in days]
    tail = [_('x.hx_required'), _('x.hx_actual'), _('x.hx_missing'), _('x.hx_ot'),
            _('x.hx_total'), _('x.hx_late'), _('x.hx_present'), _('x.hx_absent_days')]
    if with_money:
        tail += [_('x.hx_basic'), _('x.hx_ot_value'), _('x.hx_deductions'), _('x.hx_net')]
    head += tail
    first_day_col = 7
    first_tail_col = first_day_col + len(days)
    hours_cols = set(range(first_day_col, first_tail_col)) | set(range(first_tail_col, first_tail_col + 6))
    money_cols = set(range(first_tail_col + 8, first_tail_col + 12)) if with_money else set()

    wb = Workbook()
    ws = wb.active
    ws.title = f'{month:02d}-{year}'
    ws.sheet_view.rightToLeft = True

    thin = Side(style='thin', color='BBBBBB')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    head_fill = PatternFill('solid', fgColor='1F5C4A')
    head_font = Font(bold=True, color='FFFFFF', size=9)

    ncols = len(head)
    ws.cell(row=1, column=1, value=_('x.hx_title')).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=f"{p_start.isoformat()} ← {p_end.isoformat()}"
            + (f"   ·   {dept}" if dept else '')).font = Font(size=10, color='555555')
    for r in (1, 2):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=min(ncols, 12))

    HR = 4
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=HR, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, center, border

    totals = {}
    r = HR
    for e in employees:
        dd = compute_employee_days(conn, e['id'], month, year)
        if not dd:
            continue
        m = metrics.get(e['id'], {})
        r += 1
        sick_em = sum(1 for x in dd['days']
                      if x['status'] == 'leave_sick'
                      or (x['status'] in ('leave_paid', 'leave_unpaid')
                          and x.get('leave_type') in emergency))
        req = float(m.get('required_hours') or 0)
        act = float(m.get('actual_hours') or 0)
        ot = (int(m.get('ot_weekday_mins') or 0) + int(m.get('ot_weekend_mins') or 0)
              + int(m.get('ot_holiday_mins') or 0)) / 60.0
        normal = max(0.0, act - ot)
        missing = max(0.0, req - normal)
        row = [e['employee_number'], e['arabic_name'] or e['name'], e['department'] or '',
               e['position'] or '', m.get('required_days', 0), sick_em]
        cells = []
        for x in dd['days']:
            st = x['status']
            # البصمةُ الناقصة حرفٌ لا ساعات، ولو حُسب لها جزءٌ في شفتٍ مقسّم.
            worked = (x.get('hours') or 0) > 0 and st in ('present', 'weekly_off', 'holiday')
            if worked:
                cells.append((excel_hours(x['hours']), WORKED_OFF_FILL if st in ('weekly_off', 'holiday') else None))
            else:
                code, fill, _k = DAY_CODES.get(st, ('?', None, None))
                cells.append((code, fill))
        row += [c[0] for c in cells]
        row += [excel_hours(req), excel_hours(act), excel_hours(missing), excel_hours(ot),
                excel_hours(req - missing + ot), excel_hours((m.get('late_mins') or 0) / 60.0),
                m.get('actual_days', 0), m.get('absent_days', 0)]
        if with_money:
            pr = money.get(e['id']) or {}
            ot_val = sum(a['amount'] for a in pr.get('allowances', []) if a['code'] == 'overtime')
            row += [pr.get('basic', 0), ot_val, pr.get('total_deductions', 0), pr.get('net', 0)]

        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = border
            c.alignment = center if i != 2 else Alignment(vertical='center')
            if first_day_col <= i < first_tail_col:
                fill = cells[i - first_day_col][1]
                if fill:
                    c.fill = PatternFill('solid', fgColor=fill)
            if i in hours_cols and isinstance(v, float):
                c.number_format = EXCEL_HOURS_FORMAT
            if i in money_cols:
                c.number_format = MONEY_FORMAT
            if i >= first_tail_col and isinstance(v, (int, float)):
                totals[i] = totals.get(i, 0) + v

    # المجاميع
    if r > HR:
        r += 1
        ws.cell(row=r, column=2, value=_('x.hx_totals')).font = Font(bold=True)
        for i, v in totals.items():
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border, c.alignment = Font(bold=True), border, center
            if i in hours_cols:
                c.number_format = EXCEL_HOURS_FORMAT
            if i in money_cols:
                c.number_format = MONEY_FORMAT

    # مفتاحُ الحروف
    r += 2
    ws.cell(row=r, column=2, value=_('x.hx_legend')).font = Font(bold=True)
    seen = set()
    for st, (code, fill, key) in DAY_CODES.items():
        if code in seen:
            continue
        seen.add(code)
        r += 1
        c = ws.cell(row=r, column=1, value=code)
        c.alignment, c.border = center, border
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)
        ws.cell(row=r, column=2, value=_(key))
    r += 1
    c = ws.cell(row=r, column=1, value='8:00')
    c.fill, c.alignment, c.border = PatternFill('solid', fgColor=WORKED_OFF_FILL), center, border
    ws.cell(row=r, column=2, value=_('x.hx_worked_off'))

    widths = {1: 9, 2: 30, 3: 16, 4: 18, 5: 8, 6: 9}
    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(
            i, 7.5 if first_day_col <= i < first_tail_col else 11)
    ws.row_dimensions[HR].height = 32
    ws.freeze_panes = ws.cell(row=HR + 1, column=first_day_col)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
